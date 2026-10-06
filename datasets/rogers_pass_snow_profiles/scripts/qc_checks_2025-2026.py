"""Quality-control checks for the 2025-2026 period of rogers_pass_snow_profiles.

Read-only. Prints the evidence behind each issue in artifacts/2025-2026/qc_report.md.

    PYTHONIOENCODING=utf-8 uv run python datasets/rogers_pass_snow_profiles/scripts/qc_checks_2025-2026.py
"""

import hashlib
import re
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

DATASET = Path(__file__).resolve().parent.parent
RAW = DATASET / "raw_data" / "2025-2026"
V1 = DATASET / "published" / "v1"
TEMPLATE = RAW / "templates" / "StratiTemplate.xlsx"


def find(name: str) -> Path:
    """Return the unique raw file with this name.

    Parameters
    ----------
    name : str
        File name.

    Returns
    -------
    Path
        Path under raw_data/2025-2026.
    """
    hits = [p for p in RAW.rglob(name) if "templates" not in p.parts]
    assert len(hits) == 1, (name, hits)
    return hits[0]


def rows(path: Path, sheet: str) -> list[tuple]:
    """Return non-empty rows of a sheet, trailing empty cells removed.

    Parameters
    ----------
    path : Path
        Workbook path.
    sheet : str
        Sheet name.

    Returns
    -------
    list[tuple]
        Rows with at least one non-empty cell.
    """
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet]
    out = []
    for r in ws.iter_rows(values_only=True):
        vals = list(r)
        while vals and vals[-1] in (None, ""):
            vals.pop()
        if vals:
            out.append(tuple(vals))
    return out


def strat_workbooks() -> list[Path]:
    """Return stratigraphy workbooks (not spatial linkage, not template)."""
    return sorted(p for p in RAW.rglob("*.xlsx") if "templates" not in p.parts and "spatial" not in p.name.lower())


def section(title: str) -> None:
    """Print a section header."""
    print(f"\n{'=' * 8} {title}")


# ── A1: duplicated stability-test block ─────────────────────────────────
section("A1 data sheets duplicated between workbooks")
tmpl = {ws: rows(TEMPLATE, ws) for ws in openpyxl.load_workbook(TEMPLATE, read_only=True).sheetnames}
st = {p.name: rows(p, "Stability Tests")[1:] for p in strat_workbooks()}
for sheet in ("Stability Tests", "BHG", "Density", "IRIS"):
    by_block = defaultdict(list)
    for p in strat_workbooks():
        if sheet not in openpyxl.load_workbook(p, read_only=True).sheetnames:
            continue
        block = tuple(rows(p, sheet)[1:])
        if block and block != tuple(tmpl.get(sheet, [])[1:]):   # ignore unfilled templates
            by_block[block].append(p.name)
    for block, names in by_block.items():
        if len(names) > 1:
            print(f"{sheet}: identical {len(block)}-row block in {names}")
            for r in block[:8]:
                print("   ", r[:5])

# ── A2/A5: dual-Ku linkage sheets vs files ─────────────────────────────
section("A2/A5 dual-Ku linkage sheets vs files on disk")
dku_files = defaultdict(set)
for p in RAW.rglob("*.txt"):
    if "dku" in p.parts and p.name[:4].isdigit():
        dku_files[p.relative_to(RAW).parts[0]].add(int(p.name[:4]))
for p in strat_workbooks():
    wb = openpyxl.load_workbook(p, read_only=True)
    for sheet in ("dku", "Dual KU"):
        if sheet not in wb.sheetnames:
            continue
        data = rows(p, sheet)
        pairs = [(r[0], r[1]) for r in data[1:] if len(r) > 1 and isinstance(r[0], (int, float)) and isinstance(r[1], (int, float))]
        if not pairs:
            print(f"{p.name}/{sheet}: no angle-counter pairs (empty template)")
            continue
        day = p.relative_to(RAW).parts[0]
        refs = sorted({int(c) for _, c in pairs})
        print(f"{p.name}/{sheet}: angles {sorted({int(a) for a, _ in pairs})} -> counters {refs[0]}-{refs[-1]} "
              f"| missing on disk: {[c for c in refs if c not in dku_files[day]]} | files on day: {min(dku_files[day])}-{max(dku_files[day])}")
for day, nums in sorted(dku_files.items()):
    print(f"{day}: dual-Ku counters on disk {min(nums)}-{max(nums)}")

# ── A3: workbook identical to blank template ───────────────────────────
section("A3 workbooks vs blank template")
for p in strat_workbooks():
    wb = openpyxl.load_workbook(p, read_only=True)
    same = [s for s in wb.sheetnames if s in tmpl and rows(p, s) == tmpl[s]]
    diff = [s for s in wb.sheetnames if s not in same]
    if len(diff) <= 2:
        print(f"{p.name}: sheets equal to template {same}; differing {diff}")
        for s in diff:
            extra = [r for r in rows(p, s) if r not in tmpl.get(s, [])]
            print(f"    {s}: {len(extra)} rows not in template, e.g. {extra[:3]}")

# ── A4: SMP linkage vs files ───────────────────────────────────────────
section("A4 SMP references vs .pnt files")
smp_files = defaultdict(set)
for p in RAW.rglob("*.pnt"):
    smp_files[p.relative_to(RAW).parts[0]].add(int(p.stem[-4:]))
for day, nums in sorted(smp_files.items()):
    print(f"{day}: SMP files {min(nums)}-{max(nums)} ({len(nums)})")
for p in strat_workbooks():
    if "smp" not in openpyxl.load_workbook(p, read_only=True).sheetnames:
        continue
    day = p.relative_to(RAW).parts[0]
    refs = sorted({int(v) for r in rows(p, "smp")[1:] for v in r if isinstance(v, (int, float)) and v >= 100})
    if refs:
        print(f"{p.name}: refs {refs} | missing: {[r for r in refs if r not in smp_files[day]]}")
for p in sorted(RAW.rglob("*.xlsx")):
    if "spatial" in p.name.lower():
        data = rows(p, openpyxl.load_workbook(p, read_only=True).sheetnames[0])
        head = [str(h).strip() if h else "" for h in data[0]]
        if "SMP" in head:
            i = head.index("SMP")
            refs = [int(r[i]) for r in data[1:] if len(r) > i and isinstance(r[i], (int, float))]
            day = p.relative_to(RAW).parts[0]
            print(f"{p.name}: SMP refs {min(refs)}-{max(refs)} | missing: {[r for r in refs if r not in smp_files[day]]}")

# ── B: header and layout checks ────────────────────────────────────────
section("B1 AVY profile layer-table header row")
for p in strat_workbooks():
    hdr = next((r for r in rows(p, "AVY profile") if r and str(r[0]).strip().upper() == "HEIGHT"), None)
    print(f"{p.name}: {hdr}")

section("B3 IRIS sheet field-notes boilerplate")
tmpl_iris = {str(v) for r in tmpl["IRIS"] for v in r if isinstance(v, str) and len(v) > 40}
for p in strat_workbooks():
    long_text = {str(v) for r in rows(p, "IRIS") for v in r if isinstance(v, str) and len(v) > 40}
    print(f"{p.name}: template boilerplate present: {bool(long_text & tmpl_iris)}; other notes: {len(long_text - tmpl_iris)}")

# ── C: vocabulary ──────────────────────────────────────────────────────
def layer_values(paths: list[Path], col: str) -> Counter:
    """Count values in one AVY-profile layer column across workbooks."""
    c = Counter()
    for p in paths:
        data = rows(p, "AVY profile")
        hdr_i = next((i for i, r in enumerate(data) if r and str(r[0]).strip().upper() == "HEIGHT"), None)
        if hdr_i is None or col not in data[hdr_i]:
            continue
        j = data[hdr_i].index(col)
        c.update(str(r[j]).strip() for r in data[hdr_i + 1:] if len(r) > j and r[j] not in (None, ""))
    return c


v1_strat = sorted((V1 / "snow_stratigraphy").glob("*.xlsx"))
for col in ("RESISTANCE", "FORM"):
    section(f"C {col} vocabulary (2025-2026 vs v1)")
    new, old = layer_values(strat_workbooks(), col), layer_values(v1_strat, col)
    print("2025-2026:", dict(new.most_common()))
    print("v1:       ", dict(old.most_common()))
    print("only in 2025-2026:", sorted(set(new) - set(old)))

section("C Stability Tests vocabulary")
tests, scores, fracture = Counter(), Counter(), Counter()
for block in st.values():
    for r in block:
        r = list(r) + [None] * 5
        tests[str(r[0]).strip()] += 1
        scores[str(r[1]).strip()] += 1
        fracture[str(r[4]).strip()] += 1
print("Test:", dict(tests)); print("Score:", dict(scores)); print("Fracture char.:", dict(fracture))
v1_tests = Counter(str(r[0]).strip() for p in v1_strat for r in rows(p, "Stability Tests")[1:])
print("v1 Test:", dict(v1_tests))

section("C BHG values")
for p in strat_workbooks():
    if "BHG" not in openpyxl.load_workbook(p, read_only=True).sheetnames:
        continue
    vals = [r[1] for r in rows(p, "BHG")[1:] if len(r) > 1 and r[1] is not None]
    kinds = Counter(type(v).__name__ for v in vals)
    text = [v for v in vals if isinstance(v, str)][:4]
    print(f"{p.name}: {len(vals)} values {dict(kinds)} e.g. {text}")

# ── D: dates, names, sites ─────────────────────────────────────────────
section("D header DATE / LOCATION per workbook")
for p in strat_workbooks():
    data = rows(p, "AVY profile")[:4]
    flat = [v for r in data for v in r]
    def after(label):
        for i, v in enumerate(flat):
            if isinstance(v, str) and v.strip().upper() == label:
                return next((x for x in flat[i + 1:i + 4] if x not in (None, "")), None)
    print(f"{p.relative_to(RAW).parts[0][:4]} {p.name}: DATE={after('DATE')!r} LOCATION={after('LOCATION')!r} TITLE={flat[2] if len(flat) > 2 else None!r}")

section("D site spellings in file and folder names")
names = Counter()
for p in RAW.rglob("*"):
    for part in p.relative_to(RAW).parts:
        for m in re.findall(r"(hermit\w*|jim\s?bay\w*|gopher\w*|round\s?hill\w*|fidelity\w*|rogers?\W?s?\s?pass)", part, flags=re.I):
            names[m] += 1
print(dict(names.most_common()))

section("D K-band subfolder dates vs day folder")
for p in sorted({p.parent for p in RAW.rglob("*.txt") if "Radar_K" in p.parts or "ka" in p.parts}):
    print(f"{p.relative_to(RAW).parts[0]} -> {p.name}")

section("E extension casing")
print(Counter(p.suffix for p in RAW.rglob("*") if p.is_file()))

# ── F: duplicates and short files ──────────────────────────────────────
section("F duplicate files (same content)")
hashes = defaultdict(list)
for p in RAW.rglob("*"):
    if p.is_file() and p.stat().st_size > 0:
        hashes[hashlib.sha256(p.read_bytes()).hexdigest()].append(p.relative_to(RAW).as_posix())
for h, ps in hashes.items():
    if len(ps) > 1:
        print(ps)

section("F short or empty sensor files")
for p in sorted(RAW.rglob("*.txt")):
    n = len(p.read_text(encoding="utf-8", errors="replace").splitlines()) if p.stat().st_size else 0
    if n < 20:
        print(f"{p.relative_to(RAW).as_posix()}: {n} lines, {p.stat().st_size} bytes")

# ── G: physical ranges ─────────────────────────────────────────────────
section("G density range and placeholder zeros")
dens = []
for p in strat_workbooks():
    if "Density" in openpyxl.load_workbook(p, read_only=True).sheetnames:
        dens += [(p.name, r[2]) for r in rows(p, "Density")[1:] if len(r) > 2 and isinstance(r[2], (int, float))]
real = [v for _, v in dens if v > 0]
print(f"density > 0: n={len(real)} min={min(real)} max={max(real)} kg/m3")
print("zero densities (formula on empty weight) per workbook:", dict(Counter(n for n, v in dens if v == 0)))
