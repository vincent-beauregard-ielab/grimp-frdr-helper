"""Build the 2025-2026 data exploration notebook for rogers_pass_snow_profiles.

Writes notebooks/data_exploration_2025-2026.ipynb. Execute it afterwards with:

    uv run jupyter nbconvert --to notebook --execute --inplace \
        datasets/rogers_pass_snow_profiles/notebooks/data_exploration_2025-2026.ipynb

The notebook reads raw_data/2025-2026/ (new period) and published/v1/ (baseline
of the previous version) and never writes to either.
"""

from pathlib import Path

import nbformat as nbf

DATASET_ROOT = Path(__file__).resolve().parent.parent
NB_PATH = DATASET_ROOT / "notebooks" / "data_exploration_2025-2026.ipynb"

cells: list = []


def md(source: str) -> None:
    """Append a markdown cell.

    Parameters
    ----------
    source : str
        Markdown text.
    """
    cells.append(nbf.v4.new_markdown_cell(source.strip()))


def code(source: str) -> None:
    """Append a code cell.

    Parameters
    ----------
    source : str
        Python source.
    """
    cells.append(nbf.v4.new_code_cell(source.strip()))


md("""
# Data exploration — Rogers Pass snow profiles, period 2025-2026

Extension of the published dataset (v1, DOI 10.20383/103.01634) with the March 2026 campaign
(2026-03-03 to 2026-03-10). Read-only inspection of `raw_data/2025-2026/`, plus a cross-period
comparison with the published v1 files in `published/v1/`.

Summary and findings: `artifacts/2025-2026/data_exploration.md`.
""")

code(r'''
import re
import zipfile
import warnings
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd
from snowmicropyn import Profile

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
pd.set_option("display.width", 200)
pd.set_option("display.max_columns", 30)
pd.set_option("display.max_colwidth", 60)

DATASET = Path("..").resolve()          # notebook runs from notebooks/
RAW = DATASET / "raw_data" / "2025-2026"
V1 = DATASET / "published" / "v1"
assert RAW.is_dir() and V1.is_dir(), "raw_data/2025-2026 or published/v1 missing"

DAY_DATES = {
    "Day1_RoundHill_03-03-2026": "2026-03-03",
    "Day2_Fidelity_RH_04-03-2026": "2026-03-04",
    "DAY3_RoundHill Drone and Rogers Pass_05-03-2026": "2026-03-05",
    "DAY4_JimBay-RoundHill-GopherButte_20260306": "2026-03-06",
    "DAY5_RestDay_20260307": "2026-03-07",
    "DAY6_ReposeEtFilm_20260308": "2026-03-08",
    "DAY7_HermitMeadows_20260309": "2026-03-09",
    "DAY8_Fidelity_20260310": "2026-03-10",
}


def rel(p: Path) -> str:
    return p.relative_to(RAW).as_posix()


def day_of(p: Path) -> str:
    return p.relative_to(RAW).parts[0]
''')

# ── 1. Inventory ──────────────────────────────────────────────────────
md("""
## 1. File inventory

Every file is classified by instrument / content class. Classes follow the scope agreed at Step 1:
out-of-scope classes are hazard assessment PDFs, field notebook photos, logistics and blank templates.
""")

code(r'''
def classify(p: Path) -> str:
    parts = [s.lower() for s in p.relative_to(RAW).parts]
    name, ext = p.name.lower(), p.suffix.lower()
    if parts[0] == "templates":
        return "template (excluded)"
    if ext in (".heic", ".jpg", ".jpeg", ".png"):
        return "photo (excluded)"
    if ext == ".docx":
        return "field readme / notes"
    if ext == ".pdf":
        return "hazard assessment PDF (excluded)"
    if "dku" in parts:
        return "radar dual-Ku"
    if "radar_k" in parts or "ka" in parts:
        return "radar K-band"
    if "smp" in parts and ext == ".pnt":
        return "SMP"
    if "snowscope" in parts and ext == ".csv":
        return "SnowScope"
    if "iris" in parts or ext == ".txt" and "iris" in name:
        return "IRIS"
    if ext in (".kml", ".kmz") or "gps" in name or any("gps" in s for s in parts[1:-1]) and ext == ".csv":
        return "GNSS"
    if ext == ".xlsx" and "spatial" in name:
        return "spatial linkage workbook"
    if ext == ".xlsx":
        return "stratigraphy workbook"
    if ext == ".docx":
        return "field readme / notes"
    return "other"


files = sorted(p for p in RAW.rglob("*") if p.is_file())
inv = pd.DataFrame({
    "path": [rel(p) for p in files],
    "day": [day_of(p) for p in files],
    "ext": [p.suffix.lower() for p in files],
    "size": [p.stat().st_size for p in files],
    "cls": [classify(p) for p in files],
})
print(len(inv), "files")
print(inv.groupby("cls").agg(files=("path", "size"), bytes=("size", "sum")).sort_values("files", ascending=False))
''')

code(r'''
print(pd.crosstab(inv["day"], inv["cls"]).to_string())
print("\nUnclassified:", inv.loc[inv.cls == "other", "path"].tolist())
print("Empty files:", inv.loc[inv["size"] == 0, "path"].tolist())
''')

# ── 2. Stratigraphy workbooks ─────────────────────────────────────────
md("""
## 2. Stratigraphy workbooks (StratiTemplate)

Sheets present, `AVY profile` header fields (date, time, elevation, aspect, incline, location, observer),
and filled data rows per sheet. v1 workbooks had only `AVY profile`, `Stability Tests`, `Density`, `IRIS`.
""")

code(r'''
def nonempty_rows(ws, start=1):
    return [r for r in ws.iter_rows(min_row=start, values_only=True) if any(v not in (None, "") for v in r)]


def header_fields(ws) -> dict:
    """Return label -> value pairs found in the first 6 rows of an AVY profile sheet."""
    out = {}
    for row in ws.iter_rows(max_row=6, values_only=True):
        vals = list(row)
        for i, v in enumerate(vals):
            if isinstance(v, str) and v.strip().upper() in {"DATE", "TIME", "ELEVATION", "ASPECT", "INCLINE", "LOCATION", "OBSERVER", "SNOW PROFILE"}:
                nxt = next((x for x in vals[i + 1:] if x not in (None, "")), None)
                out[v.strip().upper()] = nxt
    return out


strat = [p for p in files if classify(p) == "stratigraphy workbook"]
rows, sheet_rows = [], []
for p in strat:
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    h = header_fields(wb["AVY profile"]) if "AVY profile" in wb.sheetnames else {}
    rows.append({"file": p.name, "day": day_of(p)[:4], **h})
    for ws in wb.worksheets:
        sheet_rows.append({"file": p.name, "sheet": ws.title, "rows": len(nonempty_rows(ws))})
hdr = pd.DataFrame(rows)
print(len(strat), "stratigraphy workbooks")
print(hdr.to_string())
''')

code(r'''
template = openpyxl.load_workbook(RAW / "templates" / "StratiTemplate.xlsx", read_only=True, data_only=True)
template_rows = {ws.title: len(nonempty_rows(ws)) for ws in template.worksheets}
sr = pd.DataFrame(sheet_rows)
sr["template_rows"] = sr["sheet"].map(template_rows)
sr["filled"] = sr["rows"] > sr["template_rows"].fillna(0)
print(sr.pivot_table(index="file", columns="sheet", values="rows", fill_value="").to_string())
print("\nTemplate rows per sheet:", template_rows)
empty = sr.groupby("file")["filled"].sum()
print("\nWorkbooks with no sheet filled beyond the template:", empty[empty == 0].index.tolist())
''')

code(r'''
v1_strat = sorted((V1 / "snow_stratigraphy").glob("*.xlsx"))
v1_sheets = Counter(s for p in v1_strat for s in openpyxl.load_workbook(p, read_only=True).sheetnames)
new_sheets = Counter(r["sheet"] for r in sheet_rows)
print("v1 sheets:", dict(v1_sheets))
print("2025-2026 sheets:", dict(new_sheets))
print("Sheet names only in 2025-2026:", sorted(set(new_sheets) - set(v1_sheets)))
''')

code(r'''
def layer_columns(ws, max_scan=15):
    """Return the first row that looks like the layer-table header in an AVY profile sheet."""
    for row in ws.iter_rows(max_row=max_scan, values_only=True):
        vals = [str(v).strip() for v in row if v not in (None, "")]
        if any(v.upper().startswith(("HEIGHT", "DEPTH", "HS")) for v in vals) and len(vals) >= 4:
            return tuple(vals)
    return None


cols_new = Counter(layer_columns(openpyxl.load_workbook(p, read_only=True, data_only=True)["AVY profile"]) for p in strat)
cols_v1 = Counter(layer_columns(openpyxl.load_workbook(p, read_only=True, data_only=True)["AVY profile"]) for p in v1_strat)
print("2025-2026 AVY profile layer headers:")
for k, v in cols_new.items(): print(f"  {v} x {k}")
print("v1 AVY profile layer headers:")
for k, v in cols_v1.items(): print(f"  {v} x {k}")
''')

# ── 3. Linkage sheets ─────────────────────────────────────────────────
md("""
## 3. Instrument linkage (workbook sheets and spatial workbooks)

The linkage sheets and spatial workbooks map profile positions or transect points to instrument file
identifiers. Each referenced identifier is checked against the files on disk for the same day.
""")

code(r'''
def ids_on_disk(day: str) -> dict:
    """Instrument identifiers present on disk for one day folder."""
    d = RAW / day
    dku = {int(p.name[:4]) for p in d.rglob("*") if p.is_file() and "dku" in p.parts and p.name[:4].isdigit()}
    radk = {int(p.name[:4]) for p in d.rglob("*.txt") if ("Radar_K" in p.parts or "ka" in p.parts) and p.name[:4].isdigit()}
    smp = {int(p.stem[-4:]) for p in d.rglob("*.pnt")}
    ss = {int(m.group(1)) for p in d.rglob("*.csv") for m in [re.search(r"Profile(\d+)", p.name)] if m}
    return {"Dual KU": dku, "Radar K": radk, "SMP": smp, "SS": ss}


def to_int(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


link_rows = []
for p in [p for p in files if classify(p) == "spatial linkage workbook"]:
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
    data = nonempty_rows(ws)
    head = [str(v).strip() if v is not None else "" for v in data[0]]
    disk = ids_on_disk(day_of(p))
    for col in ["Dual KU", "Radar K", "SMP", "SS"]:
        if col not in head:
            continue
        i = head.index(col)
        refs = [to_int(r[i]) for r in data[1:] if i < len(r) and to_int(r[i]) is not None]
        missing = sorted(set(refs) - disk[col])
        link_rows.append({"workbook": p.name, "column": col, "refs": len(refs), "missing_on_disk": len(missing), "examples": missing[:6]})
    print(p.name, "| columns:", head, "| points:", len(data) - 1)
print()
print(pd.DataFrame(link_rows).to_string())
''')

code(r'''
sheet_link = []
for p in strat:
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    disk = ids_on_disk(day_of(p))
    for sheet, key in [("dku", "Dual KU"), ("Dual KU", "Dual KU"), ("smp", "SMP"), ("SS", "SS"), ("Snowscope", "SS"), ("SnowScope", "SS"), ("Radar K", "Radar K"), ("RadarK", "Radar K")]:
        if sheet not in wb.sheetnames:
            continue
        nums = [to_int(v) for r in nonempty_rows(wb[sheet], start=2) for v in r if to_int(v) is not None and to_int(v) >= 0]
        cand = [n for n in nums if n in disk[key] or n > 60]   # crude filter: ids, not angles/depths
        missing = sorted({n for n in cand if n not in disk[key]})
        sheet_link.append({"workbook": p.name, "sheet": sheet, "ids_found": len(set(cand)), "missing_on_disk": missing[:10]})
print(pd.DataFrame(sheet_link).to_string())
''')

# ── 4. IRIS ───────────────────────────────────────────────────────────
md("## 4. IRIS text files")

code(r'''
def read_iris(p: Path) -> pd.DataFrame:
    df = pd.read_csv(p, header=None, skipinitialspace=True)
    df.columns = ["date", "time", "label", "voltage"][: df.shape[1]] if df.shape[1] == 4 else ["time", "voltage"]
    return df


iris = []
for p in [p for p in files if classify(p) == "IRIS"]:
    df = read_iris(p)
    iris.append({"file": rel(p), "day": day_of(p)[:4], "cols": df.shape[1], "rows": len(df), "dates": sorted(df["date"].unique().tolist()) if "date" in df else None,
                 "v_min": df["voltage"].min(), "v_max": df["voltage"].max(), "labels": sorted(df["label"].unique().tolist()) if "label" in df else None})
print(pd.DataFrame(iris).to_string())
v1_iris = sorted((V1 / "iris").glob("*"))
print("\nv1 IRIS example:", v1_iris[0].name, "columns:", read_iris(v1_iris[0]).shape[1])
''')

# ── 5. SMP ────────────────────────────────────────────────────────────
md("## 5. SnowMicroPen (`.pnt`, read with `snowmicropyn`)")

code(r'''
smp = []
for p in [p for p in files if classify(p) == "SMP"]:
    prof = Profile.load(p)
    lat, lon = prof.coordinates if prof.coordinates else (None, None)
    smp.append({"file": rel(p), "day": day_of(p)[:4], "timestamp": prof.timestamp, "lat": lat, "lon": lon,
                "samples": len(prof.samples), "depth_mm": round(prof.samples.distance.max(), 1),
                "smp_serial": prof.smp_serial, "force_max_N": round(prof.samples.force.max(), 2)})
smp = pd.DataFrame(smp)
print(len(smp), "SMP profiles")
print(smp.groupby("day").agg(n=("file", "size"), first=("timestamp", "min"), last=("timestamp", "max"),
                             depth_max=("depth_mm", "max"), no_gps=("lat", lambda s: s.isna().sum())).to_string())
print("\nSerials:", smp.smp_serial.value_counts().to_dict())
print("Sample count range:", smp.samples.min(), "-", smp.samples.max())
''')

# ── 6. SnowScope ──────────────────────────────────────────────────────
md("## 6. SnowScope CSV profiles")

code(r'''
def read_snowscope(p: Path) -> tuple[dict, pd.DataFrame]:
    lines = p.read_text(encoding="utf-8", errors="replace").splitlines()
    meta, i = {}, 0
    while i < len(lines) and not lines[i].startswith("depth (mm)"):
        parts = lines[i].split(",", 1)
        if len(parts) == 2:
            meta[parts[0].strip()] = parts[1].strip()
        i += 1
    table = pd.read_csv(p, skiprows=i, na_values=["null"])
    return meta, table


ss_rows, ss_keys, ss_cols = [], Counter(), Counter()
for p in [p for p in files if classify(p) == "SnowScope"]:
    meta, t = read_snowscope(p)
    ss_keys.update(meta.keys()); ss_cols[tuple(t.columns)] += 1
    lat, lon = (meta.get("Location", ",").split(",") + [None])[:2]
    ss_rows.append({"file": p.name, "day": day_of(p)[:4], "serial": meta.get("serialNum"), "creator": meta.get("creator name"),
                    "time": meta.get("collectionTime"), "elev": pd.to_numeric(meta.get("elevation (m)"), errors="coerce"),
                    "lat": pd.to_numeric(lat, errors="coerce"), "lon": pd.to_numeric(lon, errors="coerce"),
                    "depth_mm": pd.to_numeric(meta.get("profileDepth (mm)"), errors="coerce"), "rows": len(t),
                    "hard_max": t.iloc[:, 1].max()})
ss = pd.DataFrame(ss_rows)
print(len(ss), "SnowScope profiles")
print(ss.groupby("day").agg(n=("file", "size"), serials=("serial", lambda s: sorted(s.dropna().unique())),
                            depth_min=("depth_mm", "min"), depth_max=("depth_mm", "max"), no_loc=("lat", lambda s: s.isna().sum())).to_string())
print("\nTable columns:", dict(ss_cols))
print("Creators:", ss.creator.value_counts().to_dict())
''')

code(r'''
v1_ss = sorted((V1 / "snowscope").rglob("*.csv"))
v1_keys, v1_cols = Counter(), Counter()
for p in v1_ss:
    meta, t = read_snowscope(p)
    v1_keys.update(meta.keys()); v1_cols[tuple(t.columns)] += 1
print("Header keys only in 2025-2026:", sorted(set(ss_keys) - set(v1_keys)))
print("Header keys only in v1:", sorted(set(v1_keys) - set(ss_keys)))
print("v1 table columns:", dict(v1_cols))
print("Optical reflectance column filled in 2025-2026:",
      sum(read_snowscope(p)[1].iloc[:, 2].notna().any() for p in files if classify(p) == "SnowScope" and read_snowscope(p)[1].shape[1] > 2))
''')

# ── 7. K-band radar ───────────────────────────────────────────────────
md("""
## 7. K-band FMCW radar (`Radar_K`, `ka`)

Both folder names hold files from the same unit as v1 `radar_FMCW_K/`. The radar clock was not set in 2026:
dates come from the folder names.
""")

code(r'''
def radar_k_header(p: Path) -> dict:
    h = {}
    for line in p.read_text(encoding="utf-8", errors="replace").splitlines()[:16]:
        if ":" in line and not line.startswith("#"):
            k, v = line.split(":", 1)
            h[k.strip()] = v.strip()
    return h


rk = []
for p in [p for p in files if classify(p) == "radar K-band" and p.suffix.lower() == ".txt"]:
    h = radar_k_header(p) if p.stat().st_size else {}
    rk.append({"file": rel(p), "day": day_of(p)[:4], "folder": "ka" if "ka" in p.parts else "Radar_K", "subfolder": p.parent.name,
               "size": p.stat().st_size, "radar_no": h.get("Radar No."), "date_in_file": h.get("Date"),
               "f_start": h.get("Start-Frequency [MHz]"), "f_stop": h.get("Stop-Frequency [MHz]"), "samples": h.get("Number of Samples")})
rk = pd.DataFrame(rk)
print(rk.groupby(["day", "folder", "subfolder"]).agg(n=("file", "size"), empty=("size", lambda s: (s == 0).sum()),
                                                      dates_in_file=("date_in_file", lambda s: sorted(s.dropna().unique()))).to_string())
print("\nUnits:", rk.radar_no.value_counts().to_dict(), "| bands:", rk[["f_start", "f_stop"]].drop_duplicates().values.tolist())
print("Empty files:", rk.loc[rk["size"] == 0, "file"].tolist())
v1h = radar_k_header(next((V1 / "radar_FMCW_K").rglob("*.txt")))
print("v1 unit:", v1h.get("Radar No."), v1h.get("Start-Frequency [MHz]"), v1h.get("Stop-Frequency [MHz]"))
''')

# ── 8. Dual-Ku radar ──────────────────────────────────────────────────
md("""
## 8. Dual Ku-band FMCW radar (`dku`) — new instrument

File name: `<counter>_<13|17>GHz_<site code>_<point>_<polarization>_<angle>deg.txt`.
""")

code(r'''
NAME_RE = re.compile(r"^(\d{4})_(\d{2})GHz_([A-Za-z0-9]+)_(\d+)_([HV])_(\d+)deg\.txt$")


def dku_header(p: Path) -> dict:
    h, chirps = {}, 0
    with p.open(encoding="utf-8", errors="replace") as f:
        for line in f:
            if line.startswith("# Chirp Number"):
                chirps += 1
            elif line.startswith("# ") and ":" in line and chirps == 0:
                k, v = line[2:].split(":", 1)
                h[k.strip()] = v.strip()
    h["chirps"] = chirps
    return h


dk = []
for p in [p for p in files if classify(p) == "radar dual-Ku"]:
    m = NAME_RE.match(p.name)
    h = dku_header(p) if p.suffix.lower() == ".txt" else {}
    dk.append({"file": rel(p), "day": day_of(p)[:4], "name_ok": bool(m), "counter": m and int(m.group(1)), "ghz": m and m.group(2),
               "site": m and m.group(3), "angle": m and int(m.group(6)), "pol": m and m.group(5),
               "ts": h.get("Timestamp"), "device": h.get("Device Number"), "fmin": h.get("Min Frequency"), "fmax": h.get("Max Frequency"),
               "hdr_site": h.get("Site Name"), "chirps": h.get("chirps")})
dk = pd.DataFrame(dk)
print(len(dk), "files;", (~dk.name_ok).sum(), "do not match the naming pattern:", dk.loc[~dk.name_ok, "file"].tolist())
ok = dk[dk.name_ok]
print(ok.groupby(["day", "site"]).agg(files=("file", "size"), counters=("counter", lambda s: f"{s.min()}-{s.max()}"),
                                      angles=("angle", lambda s: sorted(s.unique())), first=("ts", "min"), last=("ts", "max")).to_string())
print("\nDevices:", ok.groupby("ghz").device.unique().to_dict())
print("Bands (kHz):", ok.groupby("ghz")[["fmin", "fmax"]].first().to_dict("index"))
print("Chirps per file:", ok.chirps.value_counts().to_dict())
print("Header site != file-name site:", (ok.hdr_site != ok.site).sum())
pairs = ok.groupby(["day", "counter"]).ghz.nunique()
print("Counters without both 13 and 17 GHz files:", pairs[pairs < 2].index.tolist()[:20])
''')

# ── 9. GNSS ───────────────────────────────────────────────────────────
md("## 9. GNSS points (Emlid exports: KML, KMZ, CSV)")

code(r'''
def kml_points(text: str) -> list[tuple]:
    pts = []
    for pm in re.findall(r"<Placemark>.*?</Placemark>", text, flags=re.S):
        name = re.search(r"<name>(.*?)</name>", pm, flags=re.S)
        c = re.search(r"<coordinates>\s*([-\d.]+),([-\d.]+),?([-\d.]*)", pm)
        if c:
            pts.append((name.group(1).strip() if name else "", float(c.group(2)), float(c.group(1)), float(c.group(3) or "nan")))
    return pts


gnss = []
for p in [p for p in files if classify(p) == "GNSS"]:
    if p.suffix.lower() == ".kml":
        pts = kml_points(p.read_text(encoding="utf-8", errors="replace"))
    elif p.suffix.lower() == ".kmz":
        z = zipfile.ZipFile(p)
        pts = kml_points(z.read(next(n for n in z.namelist() if n.endswith(".kml"))).decode("utf-8", errors="replace"))
    elif p.suffix.lower() == ".csv":
        df = pd.read_csv(p)
        pts = list(zip(df["Name"].astype(str), df["Latitude"], df["Longitude"], df["Ellipsoidal height"]))
    else:
        continue
    for n, lat, lon, h in pts:
        gnss.append({"file": rel(p), "day": day_of(p)[:4], "name": n, "lat": lat, "lon": lon, "h_ellips": h})
gnss = pd.DataFrame(gnss)
print(gnss.groupby("file").agg(points=("name", "size"), lat_min=("lat", "min"), lat_max=("lat", "max"),
                               lon_min=("lon", "min"), lon_max=("lon", "max"), h_min=("h_ellips", "min"), h_max=("h_ellips", "max")).to_string())
''')

code(r'''
csvs = [p for p in files if classify(p) == "GNSS" and p.suffix.lower() == ".csv"]
for p in csvs:
    df = pd.read_csv(p)
    print(p.name, "| solution:", df["Solution status"].value_counts().to_dict(), "| correction:", df["Correction type"].value_counts().to_dict(),
          "| receiver:", df.iloc[:, -2].unique().tolist() if df.shape[1] > 2 else None)
v1_cols = list(pd.read_csv(next((V1 / "spatial_reference").glob("*gps_points.csv"))).columns)
print("\nColumns only in 2025-2026 CSV:", sorted(set(pd.read_csv(csvs[0]).columns) - set(v1_cols)))
print("Columns only in v1 CSV:", sorted(set(v1_cols) - set(pd.read_csv(csvs[0]).columns)))
''')

# ── 10. Scope summary ────────────────────────────────────────────────
md("""
## 10. Scope summary

Site table, temporal coverage and file-derived extent. Coordinates come from GNSS exports, SMP headers
and SnowScope headers (WGS84, EPSG:4326). Heights: GNSS = ellipsoidal; SnowScope `elevation (m)` = phone/app value.
""")

code(r'''
pts = pd.concat([
    gnss.assign(source="GNSS")[["day", "lat", "lon", "source"]],
    smp.dropna(subset=["lat"]).assign(source="SMP")[["day", "lat", "lon", "source"]],
    ss.dropna(subset=["lat"]).assign(source="SnowScope")[["day", "lat", "lon", "source"]],
])
# Outliers: points more than ~0.2 deg (~15-20 km) from the median of all RTK GNSS points, e.g. a phone
# location not refreshed in the SnowScope app. Listed, then excluded from the extent.
ref_lat, ref_lon = gnss.lat.median(), gnss.lon.median()
pts = pts.reset_index(drop=True)
pts["outlier"] = ((pts.lat - ref_lat).abs() > 0.2) | ((pts.lon - ref_lon).abs() > 0.2)
print("Location outliers (excluded from extent):")
print(pts[pts.outlier].groupby(["day", "source"]).agg(n=("lat", "size"), lat=("lat", "first"), lon=("lon", "first")).to_string())
print("\nSnowScope profiles per location cluster (rounded to 0.01 deg):")
print(ss.assign(cluster=ss.lat.round(2).astype(str) + ", " + ss.lon.round(2).astype(str))
        .groupby(["day", "cluster"]).agg(n=("file", "size"), elev_min=("elev", "min"), elev_max=("elev", "max")).round(1).to_string())
good = pts[~pts.outlier]
print()
print(good.groupby(["day", "source"]).agg(n=("lat", "size"), lat_min=("lat", "min"), lat_max=("lat", "max"),
                                          lon_min=("lon", "min"), lon_max=("lon", "max")).round(5).to_string())
bbox = dict(west=good.lon.min(), east=good.lon.max(), south=good.lat.min(), north=good.lat.max())
print("\nBounding box 2025-2026:", {k: round(float(v), 5) for k, v in bbox.items()})
h = gnss.h_ellips[gnss.h_ellips > 0]
print("GNSS points with height 0 (no height in export):", int((gnss.h_ellips == 0).sum()), "in", gnss.loc[gnss.h_ellips == 0, "file"].unique().tolist())
print("GNSS ellipsoidal height range (>0):", round(h.min(), 1), "-", round(h.max(), 1), "m")
ss_ok = ss.merge(good[good.source == "SnowScope"][["lat", "lon"]].drop_duplicates(), on=["lat", "lon"])
print("SnowScope app elevation range (non-outliers):", round(ss_ok.elev.min(), 1), "-", round(ss_ok.elev.max(), 1), "m")
print("Workbook ELEVATION fields:", hdr["ELEVATION"].dropna().astype(str).unique().tolist() if "ELEVATION" in hdr else None)
''')

code(r'''
dates = pd.concat(ignore_index=True, objs=[
    smp.assign(date=smp.timestamp.dt.date.astype(str), src="SMP")[["day", "date", "src"]],
    ok.assign(date=ok.ts.str[:10], src="dual-Ku")[["day", "date", "src"]],
    pd.DataFrame([{"day": r["day"], "date": d, "src": "IRIS"} for r in iris for d in (r["dates"] or [])]),
])
print(pd.crosstab([dates.day, dates.date], dates.src).to_string())
print("\nDay folders vs nominal dates:", DAY_DATES)
''')

# ── 11. Cross-period comparison ──────────────────────────────────────
md("""
## 11. Cross-period comparison with published v1

Instrument folders, formats and conventions in v1 (`published/v1/`) vs 2025-2026 raw files.
""")

code(r'''
v1_classes = {d.name: len([f for f in d.rglob("*") if f.is_file()]) for d in sorted(V1.iterdir()) if d.is_dir()}
new_classes = inv[~inv.cls.str.contains("excluded")].cls.value_counts().to_dict()
print("v1 deposit folders:", v1_classes)
print("2025-2026 in-scope classes:", new_classes)
v1_ext = Counter(f.suffix.lower() for f in V1.rglob("*") if f.is_file())
print("\nv1 extensions:", dict(v1_ext))
print("2025-2026 in-scope extensions:", inv[~inv.cls.str.contains("excluded")].ext.value_counts().to_dict())
''')

# ── 12. Synthesis ─────────────────────────────────────────────────────
md("""
## 12. Synthesis

Key counts used in `artifacts/2025-2026/data_exploration.md`.
""")

code(r'''
summary = {
    "files total": len(inv),
    "in scope (non-excluded classes)": int((~inv.cls.str.contains("excluded")).sum()),
    "stratigraphy workbooks": len(strat),
    "SMP profiles": len(smp),
    "SnowScope profiles": len(ss),
    "IRIS files": len(iris),
    "K-band radar files": len(rk),
    "K-band empty files": int((rk["size"] == 0).sum()),
    "dual-Ku files": len(dk),
    "GNSS points": len(gnss),
}
for k, v in summary.items():
    print(f"{k:35s} {v}")
''')

nb = nbf.v4.new_notebook()
nb.metadata.kernelspec = {"display_name": "Python 3", "language": "python", "name": "python3"}
nb.cells = cells
NB_PATH.parent.mkdir(parents=True, exist_ok=True)
nbf.write(nb, NB_PATH)
print(f"Wrote {NB_PATH}")
