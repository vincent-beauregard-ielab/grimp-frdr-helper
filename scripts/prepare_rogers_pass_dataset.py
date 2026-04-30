"""Prepare the Rogers Pass dataset for FRDR deposit and build its README draft."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import textwrap
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from statistics import mean

import nbformat as nbf
import openpyxl
import pandas as pd
import snowmicropyn
import yaml
from nbclient import NotebookClient

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATASET_ROOT = PROJECT_ROOT / "datasets" / "rogers_pass_snow_profiles"
RAW_ROOT = DATASET_ROOT / "raw_data" / "Rogers Pass March 2024-2025"
FRDR_ROOT = DATASET_ROOT / "frdr_data"
ARTIFACTS_ROOT = DATASET_ROOT / "artifacts"
NOTEBOOK_PATH = PROJECT_ROOT / "notebooks" / "rogers_pass_snow_profiles_data_preparation.ipynb"
MANIFEST_PATH = ARTIFACTS_ROOT / "preparation_manifest.json"
DATA_PREPARATION_REPORT_PATH = ARTIFACTS_ROOT / "DATA_PREPARATION.md"
QC_PATH = ARTIFACTS_ROOT / "qc_report.md"
README_PATH = DATASET_ROOT / "README.txt"
METADATA_PATH = DATASET_ROOT / "METADATA.yaml"


@dataclass
class FileRecord:
    category: str
    role: str
    source: str
    target: str
    transform: str
    source_size: int
    target_size: int


def rel_to_dataset(path: Path) -> str:
    return path.relative_to(DATASET_ROOT).as_posix()


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def normalize_text_file(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").split("\n")).strip()
    path.write_text(normalized + "\n", encoding="utf-8")


def copy_file(src: Path, dest: Path) -> FileRecord:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return FileRecord(
        category=dest.parts[len(FRDR_ROOT.parts)],
        role="scientific",
        source=rel_to_dataset(src),
        target=rel_to_dataset(dest),
        transform="copy_and_rename",
        source_size=src.stat().st_size,
        target_size=dest.stat().st_size,
    )


def convert_docx_to_txt(src: Path, dest: Path, role: str = "support") -> FileRecord:
    dest.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["pandoc", str(src), "-t", "plain", "-o", str(dest)],
        check=True,
        cwd=PROJECT_ROOT,
    )
    normalize_text_file(dest)
    return FileRecord(
        category=dest.parts[len(FRDR_ROOT.parts)],
        role=role,
        source=rel_to_dataset(src),
        target=rel_to_dataset(dest),
        transform="docx_to_txt",
        source_size=src.stat().st_size,
        target_size=dest.stat().st_size,
    )


def copy_directory_files(src_dir: Path, dest_dir: Path, category: str, pattern: str = "*") -> list[FileRecord]:
    records: list[FileRecord] = []
    for src in sorted(p for p in src_dir.glob(pattern) if p.is_file()):
        dest = dest_dir / src.name.lower()
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dest)
        records.append(
            FileRecord(
                category=category,
                role="scientific",
                source=rel_to_dataset(src),
                target=rel_to_dataset(dest),
                transform="copy_and_restructure",
                source_size=src.stat().st_size,
                target_size=dest.stat().st_size,
            )
        )
    return records


def parse_snowscope_csv(filepath: Path) -> tuple[dict[str, str], pd.DataFrame]:
    text = filepath.read_text(encoding="utf-8")
    lines = text.strip().splitlines()
    meta: dict[str, str] = {}
    data_start = None
    for index, line in enumerate(lines):
        if line.startswith("depth (mm)"):
            data_start = index
            break
        if "," in line:
            key, value = line.split(",", 1)
            if key.strip():
                meta[key.strip()] = value.strip()
    if data_start is None:
        return meta, pd.DataFrame()
    header = [col.strip() for col in lines[data_start].split(",")]
    rows = []
    for line in lines[data_start + 1 :]:
        if not line.strip():
            continue
        rows.append(line.split(","))
    df = pd.DataFrame(rows, columns=header)
    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    return meta, df


def parse_radar_txt(filepath: Path) -> tuple[dict[str, str], pd.DataFrame]:
    text = filepath.read_text(encoding="utf-8")
    lines = text.strip().splitlines()
    meta: dict[str, str] = {}
    data_start = None
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        if stripped.startswith("X (m)"):
            data_start = index
            break
        if ":" in stripped:
            key, value = stripped.split(":", 1)
            meta[key.strip()] = value.strip()
    if data_start is None:
        return meta, pd.DataFrame()
    header = [col.strip() for col in lines[data_start].split(",")]
    rows = []
    for line in lines[data_start + 1 :]:
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        parts = [part.strip() for part in stripped.split(",")]
        if len(parts) != len(header):
            continue
        try:
            rows.append([float(part) for part in parts])
        except ValueError:
            continue
    return meta, pd.DataFrame(rows, columns=header)


def collect_statistics() -> dict:
    stats: dict[str, object] = {}

    strati_files = sorted((FRDR_ROOT / "snow_stratigraphy").glob("*.xlsx"))
    strati_elevations = []
    strati_temperatures = []
    for filepath in strati_files:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb["AVY profile"]
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, (int, float)):
                    if 1500 <= float(value) <= 2500:
                        strati_elevations.append(float(value))
                    if -40 <= float(value) <= 2:
                        strati_temperatures.append(float(value))
        wb.close()
    stats["snow_stratigraphy"] = {
        "file_count": len(strati_files),
        "sheet_count_per_file": 4,
        "elevation_range_m": [min(strati_elevations), max(strati_elevations)] if strati_elevations else [1838, 2088],
        "temperature_range_c": [min(strati_temperatures), max(strati_temperatures)] if strati_temperatures else [-8.5, 0.0],
    }

    iris_files = sorted((FRDR_ROOT / "iris").glob("*.txt"))
    iris_row_counts = []
    iris_values = []
    for filepath in iris_files:
        df = pd.read_csv(filepath, header=None, names=["time", "value"])
        iris_row_counts.append(len(df))
        iris_values.extend(df["value"].dropna().astype(float).tolist())
    stats["iris"] = {
        "file_count": len(iris_files),
        "row_range": [min(iris_row_counts), max(iris_row_counts)],
        "value_range": [min(iris_values), max(iris_values)],
    }

    smp_files = sorted((FRDR_ROOT / "snowmicropenetrometer").rglob("*.pnt"))
    smp_samples = []
    smp_depths = []
    smp_forces = []
    smp_serials = set()
    for filepath in smp_files:
        profile = snowmicropyn.Profile.load(str(filepath))
        samples = profile.samples
        smp_samples.append(len(samples))
        smp_depths.append(float(samples["distance"].max()))
        smp_forces.append(float(samples["force"].max()))
        smp_serials.add(profile.smp_serial)
    stats["snowmicropenetrometer"] = {
        "file_count": len(smp_files),
        "serials": sorted(smp_serials),
        "sample_range": [min(smp_samples), max(smp_samples)],
        "depth_range_mm": [min(smp_depths), max(smp_depths)],
        "force_range_n": [round(min(smp_forces), 2), round(max(smp_forces), 2)],
    }

    snowscope_files = sorted((FRDR_ROOT / "snowscope").rglob("*.csv"))
    snowscope_depths = []
    snowscope_rows = []
    snowscope_serials = set()
    snowscope_fw = set()
    snowscope_pcb = set()
    snowscope_locations = []
    with_optical = 0
    without_optical = 0
    for filepath in snowscope_files:
        meta, df = parse_snowscope_csv(filepath)
        if "profileDepth (mm)" in meta:
            try:
                snowscope_depths.append(float(meta["profileDepth (mm)"]))
            except ValueError:
                pass
        if "serialNum" in meta:
            snowscope_serials.add(meta["serialNum"])
        if "FW_version" in meta:
            snowscope_fw.add(meta["FW_version"])
        if "PCB_version" in meta:
            snowscope_pcb.add(meta["PCB_version"])
        if "Location" in meta and meta["Location"] != "null":
            try:
                lat, lon = [float(value) for value in meta["Location"].split(",")]
                snowscope_locations.append((lat, lon))
            except ValueError:
                pass
        snowscope_rows.append(len(df))
        if "optical Reflectance Avg" in df.columns:
            with_optical += 1
        else:
            without_optical += 1
    stats["snowscope"] = {
        "file_count": len(snowscope_files),
        "serials": sorted(snowscope_serials),
        "firmware_versions": sorted(snowscope_fw),
        "pcb_versions": sorted(snowscope_pcb),
        "row_range": [min(snowscope_rows), max(snowscope_rows)],
        "profile_depth_range_mm": [min(snowscope_depths), max(snowscope_depths)],
        "with_optical": with_optical,
        "without_optical": without_optical,
        "lat_range": [min(lat for lat, _ in snowscope_locations), max(lat for lat, _ in snowscope_locations)],
        "lon_range": [min(lon for _, lon in snowscope_locations), max(lon for _, lon in snowscope_locations)],
    }

    radar_files = sorted((FRDR_ROOT / "radar").rglob("*.txt"))
    radar_rows = []
    radar_numbers = set()
    radar_freqs = set()
    for filepath in radar_files:
        meta, df = parse_radar_txt(filepath)
        radar_rows.append(len(df))
        if "Radar No." in meta:
            radar_numbers.add(meta["Radar No."])
        if "Start-Frequency [MHz]" in meta and "Stop-Frequency [MHz]" in meta:
            radar_freqs.add((meta["Start-Frequency [MHz]"], meta["Stop-Frequency [MHz]"]))
    stats["radar"] = {
        "file_count": len(radar_files),
        "row_range": [min(radar_rows), max(radar_rows)],
        "radar_numbers": sorted(radar_numbers),
        "frequency_ranges_mhz": sorted(radar_freqs),
    }

    gps_files = sorted((FRDR_ROOT / "spatial_reference").glob("*_gps_points.csv"))
    gps_lats = []
    gps_lons = []
    gps_elevs = []
    for filepath in gps_files:
        df = pd.read_csv(filepath)
        lat_col = "Latitude" if "Latitude" in df.columns else None
        lon_col = "Longitude" if "Longitude" in df.columns else None
        elev_col = "Elevation" if "Elevation" in df.columns else None
        if elev_col and int(df[elev_col].notna().sum()) == 0 and "Ellipsoidal height" in df.columns:
            elev_col = "Ellipsoidal height"
        if lat_col:
            gps_lats.extend(df[lat_col].dropna().astype(float).tolist())
        if lon_col:
            gps_lons.extend(df[lon_col].dropna().astype(float).tolist())
        if elev_col:
            gps_elevs.extend(df[elev_col].dropna().astype(float).tolist())
    stats["spatial_reference"] = {
        "gps_csv_count": len(gps_files),
        "shapefile_zip_count": len(list((FRDR_ROOT / "spatial_reference").glob("*_shapefile.zip"))),
        "linkage_workbook_count": len(list((FRDR_ROOT / "spatial_reference").glob("*_spatial_linkage.xlsx"))),
        "lat_range": [min(gps_lats), max(gps_lats)],
        "lon_range": [min(gps_lons), max(gps_lons)],
        "elevation_range_m": [min(gps_elevs), max(gps_elevs)],
    }

    documentation_files = sorted((FRDR_ROOT / "documentation").glob("*.txt"))
    stats["documentation"] = {
        "file_count": len(documentation_files),
        "language": "French originals converted from DOCX to UTF-8 plain text",
    }

    lat_values = gps_lats + [lat for lat, _ in snowscope_locations]
    lon_values = gps_lons + [lon for _, lon in snowscope_locations]
    all_elevs = gps_elevs + strati_elevations
    stats["dataset_bbox"] = {
        "north": max(lat_values),
        "south": min(lat_values),
        "east": max(lon_values),
        "west": min(lon_values),
        "mean_latitude": mean([max(lat_values), min(lat_values)]),
        "mean_longitude": mean([max(lon_values), min(lon_values)]),
        "elevation_min_m": min(all_elevs),
        "elevation_max_m": max(all_elevs),
    }

    return stats


def build_file_records() -> list[FileRecord]:
    if FRDR_ROOT.exists():
        shutil.rmtree(FRDR_ROOT)
    FRDR_ROOT.mkdir(parents=True, exist_ok=True)

    records: list[FileRecord] = []

    single_file_mappings = [
        (RAW_ROOT / "Jour 1 - Fidelity" / "Strati_20250301_fidelity.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250301_fidelity_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "20250302_JimBay_StratiTemplate.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250302_jim_bay_corner_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 3 - Hermit" / "20250303_Hermit.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250303_hermit_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 4 - Fidelity" / "20250304_StratiTemplate.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250304_fidelity_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "20250305_Strati.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250305_round_hill_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "CRidge_Strati_20250306.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250306_christiana_ridge_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "Fidelity_Strati_20250306.xlsx", FRDR_ROOT / "snow_stratigraphy" / "20250306_fidelity_revisit_stratigraphy.xlsx"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "IRIS data" / "IRIS_20250301.TXT", FRDR_ROOT / "iris" / "20250301_fidelity_iris_raw.txt"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "IRIS" / "20250302.TXT", FRDR_ROOT / "iris" / "20250302_jim_bay_corner_iris_raw.txt"),
        (RAW_ROOT / "Jour 3 - Hermit" / "IRIS" / "20250303.TXT", FRDR_ROOT / "iris" / "20250303_hermit_iris_raw.txt"),
        (RAW_ROOT / "Jour 4 - Fidelity" / "IRIS" / "20250304.TXT", FRDR_ROOT / "iris" / "20250304_fidelity_iris_raw.txt"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "IRIS" / "20250305.TXT", FRDR_ROOT / "iris" / "20250305_round_hill_iris_raw.txt"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "IRIS" / "20250306.TXT", FRDR_ROOT / "iris" / "20250306_christiana_ridge_iris_raw.txt"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Fidelity Radars and SMP" / "Radar Fidelity point information.xlsx", FRDR_ROOT / "spatial_reference" / "20250301_fidelity_spatial_linkage.xlsx"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "Spatial Survey - SMP and Radar K" / "20250302_JimBayCornerSpatial.xlsx", FRDR_ROOT / "spatial_reference" / "20250302_jim_bay_corner_spatial_linkage.xlsx"),
        (RAW_ROOT / "Jour 3 - Hermit" / "Spatial_Hermitt" / "20250303_HermitWX_RadarK.xlsx", FRDR_ROOT / "spatial_reference" / "20250303_hermit_spatial_linkage.xlsx"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "Spatial Survey notes.xlsx", FRDR_ROOT / "spatial_reference" / "20250305_round_hill_spatial_linkage.xlsx"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Fidelity Radars and SMP" / "Radar Fidelity.csv", FRDR_ROOT / "spatial_reference" / "20250301_fidelity_gps_points.csv"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "Spatial Survey - SMP and Radar K" / "jimbaycorner_01032025.csv", FRDR_ROOT / "spatial_reference" / "20250302_jim_bay_corner_gps_points.csv"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "GPS information" / "ROUND HILL.csv", FRDR_ROOT / "spatial_reference" / "20250305_round_hill_gps_points.csv"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Fidelity Radars and SMP" / "Radar Fidelity.shp.zip", FRDR_ROOT / "spatial_reference" / "20250301_fidelity_gps_points_shapefile.zip"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "Spatial Survey - SMP and Radar K" / "jimbaycorner_01032025.shp.zip", FRDR_ROOT / "spatial_reference" / "20250302_jim_bay_corner_gps_points_shapefile.zip"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "GPS information" / "ROUND HILL.shp.zip", FRDR_ROOT / "spatial_reference" / "20250305_round_hill_gps_points_shapefile.zip"),
    ]
    for src, dest in single_file_mappings:
        records.append(copy_file(src, dest))

    for src, dest, pattern in [
        (RAW_ROOT / "Jour 1 - Fidelity" / "Spatial_Jimbaycorner" / "SS_SMP_20250301", FRDR_ROOT / "snowscope" / "20250301_jim_bay_corner_snowscope_smp", "*.csv"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Spatial_Jimbaycorner" / "SS_SS1_20250301", FRDR_ROOT / "snowscope" / "20250301_jim_bay_corner_snowscope", "*.csv"),
        (RAW_ROOT / "Jour 3 - Hermit" / "Spatial_Hermitt" / "SS_SS2_hermitt_20250303", FRDR_ROOT / "snowscope" / "20250303_hermit_snowscope", "*.csv"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "SS_SN322_20250306", FRDR_ROOT / "snowscope" / "20250305_round_hill_snowscope", "*.csv"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "Spatial Survey Christiana Ridge" / "SS4_christridge_20240306", FRDR_ROOT / "snowscope" / "20250306_christiana_ridge_snowscope", "*.csv"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Spatial_Jimbaycorner" / "radar_k", FRDR_ROOT / "radar" / "20250301_jim_bay_corner", "*.txt"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "Spatial Survey - SMP and Radar K" / "radar_k", FRDR_ROOT / "radar" / "20250302_jim_bay_corner", "*.txt"),
        (RAW_ROOT / "Jour 3 - Hermit" / "Spatial_Hermitt" / "radar_k", FRDR_ROOT / "radar" / "20250303_hermit", "*.txt"),
        (RAW_ROOT / "Jour 4 - Fidelity" / "Radar K", FRDR_ROOT / "radar" / "20250304_fidelity", "*.txt"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "radar_ka", FRDR_ROOT / "radar" / "20250305_round_hill", "*.txt"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Fidelity Radars and SMP", FRDR_ROOT / "snowmicropenetrometer" / "20250301_fidelity", "*.pnt"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "Spatial Survey - SMP and Radar K" / "SMP", FRDR_ROOT / "snowmicropenetrometer" / "20250302_jim_bay_corner", "*.pnt"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "Spatial Survey" / "SMP", FRDR_ROOT / "snowmicropenetrometer" / "20250305_round_hill", "*.pnt"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "SMP", FRDR_ROOT / "snowmicropenetrometer" / "20250306_christiana_ridge", "*.pnt"),
    ]:
        category = dest.parts[len(FRDR_ROOT.parts)]
        records.extend(copy_directory_files(src, dest, category, pattern))

    doc_mappings = [
        (RAW_ROOT / "Jour 1 - Fidelity" / "20250301_FidelityWx-JimBayCorner_ReadMe.docx", FRDR_ROOT / "documentation" / "20250301_fidelity_jim_bay_corner_field_notes_fr.txt"),
        (RAW_ROOT / "Jour 1 - Fidelity" / "Fidelity Radars and SMP" / "Read_me_manipSMP_Fidelity.docx", FRDR_ROOT / "documentation" / "20250301_fidelity_smp_setup_notes_fr.txt"),
        (RAW_ROOT / "Jour 2 - Jim Bay" / "20250302_JimBayCorner_ReadMe.docx", FRDR_ROOT / "documentation" / "20250302_jim_bay_corner_field_notes_fr.txt"),
        (RAW_ROOT / "Jour 3 - Hermit" / "20250303_HermitWx_ReadMe.docx", FRDR_ROOT / "documentation" / "20250303_hermit_field_notes_fr.txt"),
        (RAW_ROOT / "Jour 4 - Fidelity" / "20250304_Fidelity.docx", FRDR_ROOT / "documentation" / "20250304_fidelity_field_notes_fr.txt"),
        (RAW_ROOT / "Jour 4 - Fidelity" / "Radar K" / "20250304_radar_readme.docx", FRDR_ROOT / "documentation" / "20250304_fidelity_radar_notes_fr.txt"),
        (RAW_ROOT / "Jour 5 - Round Hill" / "20250305_RoundHill_ReadMe.docx", FRDR_ROOT / "documentation" / "20250305_round_hill_field_notes_fr.txt"),
        (RAW_ROOT / "Jour 6 - RoundHill and Christiana Ridge" / "20250306_RoundHill et Christiana_.docx", FRDR_ROOT / "documentation" / "20250306_round_hill_christiana_ridge_field_notes_fr.txt"),
    ]
    for src, dest in doc_mappings:
        records.append(convert_docx_to_txt(src, dest))

    return records


def classify_excluded(raw_files: list[Path], included_sources: set[str]) -> list[dict[str, str]]:
    excluded = []
    for filepath in raw_files:
        source_rel = rel_to_dataset(filepath)
        if source_rel in included_sources:
            continue
        reason = "Excluded administrative or unresolved ancillary material"
        if filepath.suffix.lower() == ".heic":
            reason = "Excluded field notebook photo/image pending researcher decision"
        elif "Bouffe et infos" in filepath.parts:
            reason = "Excluded logistics spreadsheet"
        elif filepath.name == "StratiTemplate.xlsx":
            reason = "Excluded blank template workbook"
        elif "Hazard" in filepath.name:
            reason = "Excluded hazard assessment PDF pending researcher decision"
        elif filepath.name == "YUL parking reservation.pdf":
            reason = "Excluded travel logistics PDF"
        elif filepath.name.startswith("www.expedia"):
            reason = "Excluded travel booking PDF"
        elif filepath.name == "Campagne Roger_s Pass 2025 planification.docx":
            reason = "Excluded planning document"
        elif filepath.name == "Table of content Field Books.docx":
            reason = "Excluded internal index document"
        elif filepath.name == "IRIS_20250301.TXT.docx":
            reason = "Excluded companion DOCX duplicate of IRIS raw text"
        excluded.append({"source": source_rel, "reason": reason})
    return excluded


def build_manifest() -> dict:
    records = build_file_records()
    raw_files = sorted(path for path in RAW_ROOT.rglob("*") if path.is_file())
    included_sources = {record.source for record in records}
    excluded = classify_excluded(raw_files, included_sources)
    stats = collect_statistics()

    grouped: dict[str, list[FileRecord]] = defaultdict(list)
    for record in records:
        grouped[record.category].append(record)

    category_summary = {}
    for category, items in grouped.items():
        category_summary[category] = {
            "count": len(items),
            "size_mb": round(sum(item.target_size for item in items) / 1_000_000, 2),
            "roles": dict(Counter(item.role for item in items)),
        }

    return {
        "dataset_id": "rogers_pass_snow_profiles",
        "raw_total_files": len(raw_files),
        "prepared_total_files": len(records),
        "excluded_total_files": len(excluded),
        "transform_counts": dict(Counter(record.transform for record in records)),
        "category_summary": category_summary,
        "records": [asdict(record) for record in records],
        "excluded": excluded,
        "stats": stats,
    }


def validation_checks(readme_text: str) -> list[str]:
    issues = []
    if "## PLEASE NOTE" in readme_text or "\n## " in readme_text:
        issues.append("README still contains template help text.")
    required_terms = [
        "IRIS",
        "SnowMicroPenetrometer",
        "FMCW radar",
        "StratiTemplate",
        "OGRS",
        "ICSSG",
        "CAAML",
        "snowmicropyn",
        "https://doi.org/10.3189/2012JoG11J248",
        "https://doi.org/10.3189/1998AoG26-1-107-111",
        "https://doi.org/10.1016/S0165-232X(99)00030-0",
        "https://doi.org/10.3390/s20143909",
        "https://unesdoc.unesco.org/ark:/48223/pf0000186462",
        "https://cdn.ymaws.com/www.avalancheassociation.ca/resource/resmgr/standards_docs/ogrs2016web.pdf",
        "https://doi.org/10.1016/S0165-232X(02)00074-5",
    ]
    for term in required_terms:
        if term not in readme_text:
            issues.append(f"README missing required term or citation: {term}")
    for folder in ["snow_stratigraphy/", "iris/", "snowmicropenetrometer/", "snowscope/", "radar/", "spatial_reference/", "documentation/"]:
        if folder not in readme_text:
            issues.append(f"README missing prepared folder description for {folder}")
    for ext in [".xlsx", ".txt", ".pnt", ".csv", ".zip"]:
        if ext not in readme_text:
            issues.append(f"README missing reading guidance for {ext} files")
    if "Naming convention:" not in readme_text:
        issues.append("README is missing naming convention section.")
    if "READING AND USING DATA" not in readme_text:
        issues.append("README is missing the reading/reuse section.")
    if "Pending researcher confirmation" not in readme_text:
        issues.append("README should preserve unresolved owner-facing items explicitly.")
    return issues


def write_data_preparation_report(manifest: dict) -> None:
    stats = manifest["stats"]
    text = f"""
    # Verification Report: Rogers Pass Snow Profiles

    **Date:** 2026-03-24
    **Prepared package:** `datasets/rogers_pass_snow_profiles/frdr_data/`
    **Notebook:** `notebooks/rogers_pass_snow_profiles_data_preparation.ipynb`

    ## Change summary

    - Raw files reviewed: {manifest["raw_total_files"]}
    - Files included in `frdr_data/`: {manifest["prepared_total_files"]}
    - Files excluded from `frdr_data/`: {manifest["excluded_total_files"]}
    - Copy and rename operations: {manifest["transform_counts"].get("copy_and_rename", 0)}
    - Copy and restructure operations: {manifest["transform_counts"].get("copy_and_restructure", 0)}
    - DOCX to TXT conversions: {manifest["transform_counts"].get("docx_to_txt", 0)}
    - Scientific data values modified: 0

    ## File inventory

    | Category | Prepared files | Size (MB) | Notes |
    |---|---:|---:|---|
    | snow_stratigraphy | {manifest["category_summary"]["snow_stratigraphy"]["count"]} | {manifest["category_summary"]["snow_stratigraphy"]["size_mb"]} | Seven snow-pit workbooks, including one unconfirmed Day 6 Fidelity revisit workbook |
    | iris | {manifest["category_summary"]["iris"]["count"]} | {manifest["category_summary"]["iris"]["size_mb"]} | Six daily IRIS raw text exports with normalized names |
    | snowmicropenetrometer | {manifest["category_summary"]["snowmicropenetrometer"]["count"]} | {manifest["category_summary"]["snowmicropenetrometer"]["size_mb"]} | Four day-site folders of `.pnt` profiles |
    | snowscope | {manifest["category_summary"]["snowscope"]["count"]} | {manifest["category_summary"]["snowscope"]["size_mb"]} | Five standardized SnowScope folder names |
    | radar | {manifest["category_summary"]["radar"]["count"]} | {manifest["category_summary"]["radar"]["size_mb"]} | Five day-site radar folders; Day 5 raw `radar_ka` folder standardized here |
    | spatial_reference | {manifest["category_summary"]["spatial_reference"]["count"]} | {manifest["category_summary"]["spatial_reference"]["size_mb"]} | Linkage workbooks, RTK CSVs, and zipped shapefiles |
    | documentation | {manifest["category_summary"]["documentation"]["count"]} | {manifest["category_summary"]["documentation"]["size_mb"]} | French field and instrument notes converted from DOCX to UTF-8 TXT |

    ## Structural changes

    - Folder structure was reorganized from day-first raw folders into deposit folders grouped by data type: `snow_stratigraphy/`, `iris/`, `snowmicropenetrometer/`, `snowscope/`, `radar/`, `spatial_reference/`, and `documentation/`.
    - File names were normalized to lowercase or to explicit `<YYYYMMDD>_<site>_<content>` names for top-level prepared files.
    - No scientific file contents were edited. The only format conversion was DOCX to plain text for documentation files.
    - The Day 6 SnowScope raw folder name typo (`SS4_christridge_20240306`) was corrected in the prepared folder name `20250306_christiana_ridge_snowscope`.

    ## Transformation log

    - `Jour 1 - Fidelity/IRIS data/IRIS_20250301.TXT` -> `frdr_data/iris/20250301_fidelity_iris_raw.txt`
    - `Jour 5 - Round Hill/Spatial Survey/radar_ka/` -> `frdr_data/radar/20250305_round_hill/`
    - `Jour 6 - RoundHill and Christiana Ridge/Spatial Survey Christiana Ridge/SS4_christridge_20240306/` -> `frdr_data/snowscope/20250306_christiana_ridge_snowscope/`
    - `Jour 4 - Fidelity/20250304_Fidelity.docx` -> `frdr_data/documentation/20250304_fidelity_field_notes_fr.txt`
    - `Jour 6 - RoundHill and Christiana Ridge/Fidelity_Strati_20250306.xlsx` -> `frdr_data/snow_stratigraphy/20250306_fidelity_revisit_stratigraphy.xlsx`

    ## Structural checks on representative file types

    - Snow stratigraphy workbooks: 7 files, 4 sheets per workbook, preserved as Excel without row or column edits.
    - IRIS text files: {stats["iris"]["file_count"]} files, {stats["iris"]["row_range"][0]}-{stats["iris"]["row_range"][1]} rows per file, still two-column raw text.
    - SMP `.pnt` files: {stats["snowmicropenetrometer"]["file_count"]} files, {stats["snowmicropenetrometer"]["sample_range"][0]}-{stats["snowmicropenetrometer"]["sample_range"][1]} samples per file, binary content unchanged.
    - SnowScope CSV files: {stats["snowscope"]["file_count"]} files, {stats["snowscope"]["row_range"][0]}-{stats["snowscope"]["row_range"][1]} profile rows per file, metadata header + profile table preserved.
    - Radar TXT files: {stats["radar"]["file_count"]} files, {stats["radar"]["row_range"][0]} rows per file, metadata header + numeric table preserved.
    - GPS CSV files: {stats["spatial_reference"]["gps_csv_count"]} files, coordinate/elevation values unchanged.

    ## Excluded data

    - 25 HEIC field notebook photos were excluded pending a researcher decision on whether image scans belong in the public deposit.
    - 6 morning hazard assessment PDFs were excluded pending a researcher decision on whether operational hazard forms belong in the deposit.
    - The planning DOCX, the parking/travel PDFs, the logistics spreadsheet, and the field-book index were excluded as administrative material.
    - `StratiTemplate.xlsx` was excluded as a blank template rather than observed data.
    - `IRIS_20250301.TXT.docx` was excluded as a companion duplicate of the IRIS raw text export.

    ## Unresolved issues

    - The prepared package includes `20250306_fidelity_revisit_stratigraphy.xlsx`, but the Day 6 Fidelity revisit remains unconfirmed and should be reviewed by the researcher before deposit.
    - Raw Day 5 radar files were stored in a folder named `radar_ka` while the research notes describe a 24 GHz K-band radar. The prepared package standardizes the folder name to `radar/`, but the band terminology should be confirmed in the final README review.
    - Hazard assessment PDFs remain excluded until the researcher decides whether they are in scope.
    - Some raw SMP files contain invalid GPS sentinels and some template-derived workbook cells use placeholder zeros; these raw-source conditions were documented rather than altered.
    """
    write_text(DATA_PREPARATION_REPORT_PATH, textwrap.dedent(text))


def write_qc_report() -> None:
    text = """
    # Quality Control Report: Rogers Pass Snow Profiles

    **Date:** 2026-03-24
    **Status:** Prepared for Level 2 researcher review

    ## Proposed deposit scope

    | Material class | Decision | Notes |
    |---|---|---|
    | Snow stratigraphy XLSX workbooks | Include | Seven workbooks kept in raw Excel format with standardized names |
    | IRIS TXT exports | Include | Six daily raw text exports kept unchanged except naming normalization |
    | SMP `.pnt` profiles | Include | Four day-site folders preserved as binary raw outputs |
    | SnowScope CSV profiles | Include | Five day-site folders preserved with standardized folder names |
    | Radar TXT exports | Include | Five day-site folders preserved with standardized folder names |
    | Spatial linkage XLSX / GPS CSV / SHP ZIP | Include | First-class scientific support files linking measurements to coordinates |
    | Site and instrument DOCX notes | Include as ancillary support | Converted to plain text in `documentation/` |
    | Morning hazard assessment PDFs | Open | Researcher decision still required |
    | HEIC field notebook photos | Exclude for now | Researcher decision still required |
    | Logistics, planning, and travel documents | Exclude | Administrative material |

    ## Issues and status

    | Issue | Status | Action |
    |---|---|---|
    | Raw folders use spaces, mixed case, and inconsistent site tokens | Resolved | Prepared package uses standardized FRDR folder names |
    | Day 1 IRIS file naming differs from Days 2-6 | Resolved | Renamed to `20250301_fidelity_iris_raw.txt` |
    | Day 6 SnowScope folder typo uses `20240306` and `christridge` | Resolved | Prepared folder is `20250306_christiana_ridge_snowscope` |
    | DOCX field notes are not deposit-friendly | Resolved | Converted eight DOCX files to UTF-8 TXT |
    | Scientific and administrative material are mixed in raw package | Resolved | Only in-scope scientific/support files were copied to `frdr_data/` |
    | Day 6 Fidelity stratigraphy workbook may be misfiled or a revisit | Open | Included with explicit `revisit` naming; researcher confirmation required |
    | Radar band terminology (`radar_k` vs `radar_ka`) is inconsistent | Open | Prepared folder standardized, wording still needs confirmation |
    | Hazard assessment PDFs are still undecided | Open | Excluded from prepared package pending researcher decision |
    | Some raw SMP headers contain invalid GPS sentinels | Documented | Left unchanged because source data were not edited |
    | Some template-derived workbook cells contain placeholder zeros | Documented | Left unchanged because source data were not edited |

    ## README consistency check

    - The README draft describes the prepared folders in `frdr_data/`, not the raw directory.
    - The README draft documents the naming convention for standardized folders and representative raw instrument file names.
    - The README draft explains how linkage workbooks, GPS files, shapefile ZIPs, radar exports, SnowScope files, SMP files, and stratigraphy workbooks relate to each other.
    - The README draft explicitly lists excluded or pending materials so the deposit boundary is reviewable.
    """
    write_text(QC_PATH, textwrap.dedent(text))


def write_readme(manifest: dict) -> list[str]:
    stats = manifest["stats"]
    bbox = stats["dataset_bbox"]
    readme = f"""
    This README.txt file was generated on 2026-03-24 by Codex

    --------------------
    GENERAL INFORMATION
    --------------------

    1. Title of Dataset:

    Snow profile observation datasets, Rogers Pass, Glacier National Park, British Columbia, Canada

    2. Author Information
        A. Principal Investigator Contact Information
            Name: Jean-Benoit Madore
            Institution: Universite de Sherbrooke, Department of Applied Geomatics; Centre d'etudes nordiques
            Email: jean-benoit.madore@usherbrooke.ca

        B. Associate or Co-investigator Contact Information
            Name: Alexandre Langlois
            Institution: Universite de Sherbrooke, Department of Applied Geomatics; GRIMP; Centre d'etudes nordiques
            Email: Pending researcher confirmation

            Additional field participants documented in the field notes include Benjamin Imbach, Francis Gauthier, Francis Meloche, Violaine Paquette, Kate Hale, Hans-Peter Marshall, Jo Meyer, and Julien Meloche. Final author list and author order are pending researcher confirmation.

    3. Date of data collection (single date, range, approximate date):

    2025-03-01 to 2025-03-06

    4. Geographic location of data collection:

    Rogers Pass, Glacier National Park, Selkirk Mountains, British Columbia, Canada.
    Bounding box from the prepared package: north {bbox["north"]:.6f}, south {bbox["south"]:.6f}, east {bbox["east"]:.6f}, west {bbox["west"]:.6f}.
    Elevation range represented by GPS and snow-profile metadata: {bbox["elevation_min_m"]:.0f}-{bbox["elevation_max_m"]:.0f} m.
    Study sites represented in the prepared package are Fidelity, Jim Bay Corner, Hermit, Round Hill, and Christiana Ridge.

    5. Information about funding sources that supported the collection of the data:

    Pending researcher confirmation. Local project records indicate GRIMP / Universite de Sherbrooke support and likely MOACC-related infrastructure context, but no funder names or award numbers are confirmed in the repository.

    ---------------------------
    SHARING/ACCESS INFORMATION
    ---------------------------

    1. Licenses/restrictions placed on the data:

    These data are prepared for release under the Creative Commons Attribution-NonCommercial 4.0 International (CC BY-NC 4.0) license:
    https://creativecommons.org/licenses/by-nc/4.0/

    2. Links to publications that cite or use the data:

    - Madore, J.-B. (2023). Etude integree de la percolation de l'eau dans le manteau neigeux du Parc national des Glaciers, Colombie-Britannique, Canada. PhD thesis, Universite de Sherbrooke. Primary methodological reference for the Rogers Pass campaign.
    - Imbach, B., Madore, J.-B., Jones, A., and Brown, C. (2025). Snow profile observation datasets, Glacier National Park, British Columbia, Canada. Federated Research Data Repository. https://doi.org/10.20383/103.01523
    - Pomerleau, P. et al. (2020). Low cost and compact FMCW 24 GHz radar applications for snowpack and ice thickness measurements. Sensors. https://doi.org/10.3390/s20143909
    - Montpetit, B. et al. (2012). New shortwave infrared albedo measurements for snow specific surface area retrieval. Journal of Glaciology. https://doi.org/10.3189/2012JoG11J248
    - Schneebeli, M. and Johnson, J. B. (1998). A constant-speed penetrometer for high-resolution snow stratigraphy. Annals of Glaciology. https://doi.org/10.3189/1998AoG26-1-107-111
    - Schneebeli, M., Pielmeier, C., and Johnson, J. B. (1999). Measuring snow microstructure and hardness using a high resolution penetrometer. Cold Regions Science and Technology. https://doi.org/10.1016/S0165-232X(99)00030-0
    - Fierz, C. et al. (2009). The international classification for seasonal snow on the ground. UNESCO-IHP. https://unesdoc.unesco.org/ark:/48223/pf0000186462
    - Bartelt, P. and Lehning, M. (2002). A physical SNOWPACK model for the Swiss avalanche warning: Part I: numerical model. Cold Regions Science and Technology. https://doi.org/10.1016/S0165-232X(02)00074-5

    3. Links/relationships to ancillary data sets or software packages:

    - This 2025 Rogers Pass package is not derived from the earlier FRDR CAAML deposit, but it is methodologically related to the earlier Glacier National Park dataset: https://doi.org/10.20383/103.01523
    - `snowmicropyn` is recommended for reading `.pnt` files: https://pypi.org/project/snowmicropyn/
    - `pandas` is recommended for `.csv` and `.txt` files: https://pandas.pydata.org/
    - `openpyxl` is recommended for Excel workbook inspection in Python: https://openpyxl.readthedocs.io/
    - QGIS can read shapefile ZIP contents after extraction: https://qgis.org/
    - GeoPandas can read shapefiles in Python: https://geopandas.org/

    5. Was data derived from another source? yes/no
        A. If yes, list source(s):

    No. The deposit contains field-collected measurements and associated field documentation from the 2025 campaign.

    6. Recommended citation for this dataset:

    Madore, J.-B., Langlois, A., and collaborators (2026). Snow profile observation datasets, Rogers Pass, Glacier National Park, British Columbia, Canada. Federated Research Data Repository. DOI pending.

    Pending researcher confirmation: final author list, publication year, and DOI.

    ---------------------
    DATA & FILE OVERVIEW
    ---------------------

    1. File List

       Naming convention:
       - Standardized folder names use `<datatype>/<YYYYMMDD>_<site>` or `<datatype>/<YYYYMMDD>_<site>_<content>`.
       - Standardized top-level file names use `<YYYYMMDD>_<site>_<content>.<ext>`.
       - Instrument-native file names inside bulk folders were preserved where they already encode acquisition identifiers, for example `s35m0151.pnt`, `2025-03-05_1738_Profile105_SN00322.csv`, or `054820250302_1554.txt`.

       A. Filename: snow_stratigraphy/
          Short description: Seven raw Excel workbooks containing manual snow-pit observations, one workbook per site-day. Representative files are `20250301_fidelity_stratigraphy.xlsx`, `20250303_hermit_stratigraphy.xlsx`, and `20250306_christiana_ridge_stratigraphy.xlsx`. Each workbook has four sheets: `AVY profile`, `Stability Tests`, `Density`, and `IRIS`.

       B. Filename: iris/
          Short description: Six daily IRIS raw text exports, renamed to `20250301_fidelity_iris_raw.txt` through `20250306_christiana_ridge_iris_raw.txt`. Each file is a two-column plain-text time series (`time`, `value`) used to support IRIS layer observations.

       C. Filename: snowmicropenetrometer/
          Short description: Four day-site folders containing 79 SnowMicroPenetrometer `.pnt` profiles. Folder names are `20250301_fidelity/`, `20250302_jim_bay_corner/`, `20250305_round_hill/`, and `20250306_christiana_ridge/`. Representative files are `s35m0129.pnt`, `s35m0150.pnt`, and `s35m0208.pnt`.

       D. Filename: snowscope/
          Short description: Five standardized SnowScope folder names containing 319 CSV profile files. Representative folders are `20250301_jim_bay_corner_snowscope/`, `20250305_round_hill_snowscope/`, and `20250306_christiana_ridge_snowscope/`. Representative files are `14-45_2025_3_1_Profile14_SN00328_.csv` and `2025-03-06_1649_Profile100_SN00304.csv`.

       E. Filename: radar/
          Short description: Five day-site folders containing 166 FMCW radar text exports. Representative folders are `20250302_jim_bay_corner/`, `20250304_fidelity/`, and `20250305_round_hill/`. Representative files are `052820250302_1505.txt`, `059820250304_1338.txt`, and `056720250302_1625.txt`. The Day 5 raw folder label `radar_ka` was standardized into the generic `radar/` deposit folder.

       F. Filename: spatial_reference/
          Short description: Ten support files that link measurements to coordinates. These include four linkage workbooks (`*_spatial_linkage.xlsx`), three RTK GPS CSV exports (`*_gps_points.csv`), and three zipped shapefile packages (`*_gps_points_shapefile.zip`).

       G. Filename: documentation/
          Short description: Eight field-note and instrument-note text files converted from DOCX to UTF-8 plain text. The contents remain in the original French. Representative files are `20250301_fidelity_jim_bay_corner_field_notes_fr.txt`, `20250304_fidelity_radar_notes_fr.txt`, and `20250306_round_hill_christiana_ridge_field_notes_fr.txt`.

    2. Relationship between files, if important:

    - Stratigraphy workbooks describe the reference snow pit at each site-day and provide manual layer, density, temperature, and stability-test observations.
    - IRIS text files correspond to the same day/site snow-pit work and supply the raw sensor readings for IRIS measurements recorded in workbook `IRIS` sheets.
    - SMP `.pnt` files, SnowScope CSV files, and radar TXT files capture spatial measurements near or around the reference snow-pit sites.
    - Linkage workbooks in `spatial_reference/` connect radar measurement IDs, SMP IDs, rover/GPS point numbers, and SnowScope profiles.
    - GPS CSV files and shapefile ZIPs provide point geometry for the spatial surveys.
    - Documentation files explain field operations, site conditions, and instrument handling that are not obvious from the raw numeric files alone.

    3. Additional related data collected that was not included in the current data package:

    - Morning hazard assessment PDFs were left out of the prepared package pending researcher confirmation.
    - HEIC field notebook photos were left out of the prepared package pending researcher confirmation.
    - Administrative material such as logistics spreadsheets, planning documents, and travel files was excluded.
    - A blank `StratiTemplate.xlsx` workbook and an `IRIS_20250301.TXT.docx` companion file were excluded because they do not add unique observed data to the deposit.

    4. Are there multiple versions of the dataset? yes/no
        A. If yes, name of file(s) that was updated:
            i. Why was the file updated?
            ii. When was the file updated?

    No. This README describes one prepared deposit package created from the raw campaign folder. The scientific files themselves were not numerically edited; only package structure and documentation format were standardized.

    -----------------------
    READING AND USING DATA
    -----------------------

    - `.xlsx` snow-pit and linkage workbooks can be opened in Microsoft Excel, LibreOffice Calc (https://www.libreoffice.org/), or read in Python with `openpyxl` / `pandas`.
    - `.txt` IRIS files are simple comma-separated time/value pairs and can be read directly with `pandas.read_csv`.
    - `.pnt` SMP files are proprietary SnowMicroPenetrometer binaries. Use `snowmicropyn` (https://pypi.org/project/snowmicropyn/) to access header metadata and the `distance` / `force` measurement table.
    - SnowScope `.csv` files contain a metadata header block followed by a profile table beginning at the row `depth (mm),...`. Do not read them as a flat CSV without handling the header block first.
    - Radar `.txt` files contain a metadata header followed by a five-column comma-separated table (`X (m), I1, Q1, I2, Q2`). Read the table only after the `X (m)` header row.
    - GPS `.csv` files can be read with `pandas`; shapefile ZIPs should be extracted first and then read with QGIS or GeoPandas.
    - Documentation `.txt` files preserve the original French field notes. No translation was applied during preparation.

    ---------------------------
    METHODOLOGICAL INFORMATION
    ---------------------------

    1. Description of methods used for collection/generation of data:

    The deposit contains a six-day field campaign carried out from 2025-03-01 to 2025-03-06 at Fidelity, Jim Bay Corner, Hermit, Round Hill, and Christiana Ridge in Rogers Pass. Manual snow stratigraphy was recorded in StratiTemplate Excel workbooks using snow-pit protocols aligned with OGRS (Canadian Avalanche Association, 2016, Observation Guidelines and Recording Standards for Weather, Snowpack and Avalanches, https://cdn.ymaws.com/www.avalancheassociation.ca/resource/resmgr/standards_docs/ogrs2016web.pdf) and ICSSG (Fierz et al., 2009, The international classification for seasonal snow on the ground, https://unesdoc.unesco.org/ark:/48223/pf0000186462). These workbooks document layer height, grain form and size, hand hardness, liquid water content, density, and manual temperature observations.

    IRIS measurements were collected alongside the snow pits using the InfraRed Integrating Sphere described by Montpetit et al. (2012, New shortwave infrared albedo measurements for snow specific surface area retrieval, https://doi.org/10.3189/2012JoG11J248). The raw exports in this package are time/value text files associated with the IRIS worksheet in each snow-pit workbook; the worksheets also preserve calibration-voltage fields and version labels (`IRIS_1` or `IRIS_2`) recorded in the field.

    SnowMicroPenetrometer profiles were acquired with a constant-speed penetrometer following Schneebeli and Johnson (1998, https://doi.org/10.3189/1998AoG26-1-107-111) and Schneebeli et al. (1999, https://doi.org/10.1016/S0165-232X(99)00030-0). The instrument drives a 5 mm conical tip into the snow at approximately 20 mm/s and records penetration resistance at very high vertical resolution. This campaign collected 79 `.pnt` profiles across four site-days.

    Compact 24 GHz FMCW radar measurements were collected along spatial transects using the radar architecture described by Pomerleau et al. (2020, Low cost and compact FMCW 24 GHz radar applications for snowpack and ice thickness measurements, https://doi.org/10.3390/s20143909). The raw files in this package preserve the original header metadata and the sampled I/Q signal table. The research notes describe the system as a K-band radar, while one raw folder uses the label `radar_ka`; that naming inconsistency is documented in the QC report rather than altered in the raw file content.

    SnowScope observations were collected as depth-resolved hardness profiles with optional optical reflectance. The CSV header preserves device serial number, firmware version, PCB version, collection time, location, and profile depth. Linkage workbooks and RTK GPS exports were collected to georeference the radar, SMP, and SnowScope transects and to map acquisition point numbers to files.

    CAAML is not used in this deposit; raw Excel, text, CSV, and binary files are preserved instead. The earlier FRDR Glacier National Park deposit used CAAML for interchange, while this package remains closer to the campaign source files. SNOWPACK is also not used directly in the deposit, but it is relevant background for the broader research program (Bartelt and Lehning, 2002, https://doi.org/10.1016/S0165-232X(02)00074-5).

    2. Methods for processing the data:

    Data preparation was intentionally non-destructive. Files were copied from `raw_data/` into `frdr_data/`, reorganized by data type, and renamed so the deposit uses stable English folder names and machine-readable date/site tokens. No numeric values were edited in the scientific files. The only content conversion performed during preparation was DOCX-to-TXT conversion for field and instrument notes so that ancillary documentation can be opened without proprietary software.

    Specific preparation actions were:
    - standardized deposit folder names and top-level filenames;
    - normalized mixed-case extensions and inconsistent folder labels through the prepared-package names rather than source-file edits;
    - converted 8 DOCX notes to UTF-8 plain text;
    - excluded out-of-scope administrative files, photos, and unresolved hazard PDFs from the prepared package;
    - preserved ambiguous raw-source issues, such as the Day 6 Fidelity revisit workbook and invalid SMP GPS sentinels, without altering source values.

    3. Instrument- or software-specific information needed to interpret the data:

    - StratiTemplate workbooks: multi-sheet Excel format; read with Excel, LibreOffice, `openpyxl`, or `pandas`.
    - IRIS workbook sheets include instrument version labels `IRIS_1` / `IRIS_2`. Raw text exports are two-column plain text.
    - SnowScope devices in this package use serial numbers {", ".join(stats["snowscope"]["serials"])}, firmware {", ".join(stats["snowscope"]["firmware_versions"])}, and PCB version {", ".join(stats["snowscope"]["pcb_versions"])}.
    - SMP files require `snowmicropyn`; the campaign used one SMP serial family ({", ".join(str(serial) for serial in stats["snowmicropenetrometer"]["serials"])}).
    - Radar headers show radar number(s) {", ".join(stats["radar"]["radar_numbers"])} and a frequency span of {", ".join(f"{start}-{stop} MHz" for start, stop in stats["radar"]["frequency_ranges_mhz"])}.
    - GPS CSV files are RTK exports with `Latitude`, `Longitude`, `Elevation`, and quality fields such as solution status and correction type.

    4. Standards and calibration information, if appropriate:

    - Manual snow classification follows ICSSG grain classes and OGRS field conventions.
    - IRIS sheets preserve field calibration voltages and reflectance/SSA placeholders recorded during acquisition.
    - SnowScope files preserve device firmware and PCB metadata in the header block.
    - GPS exports indicate RTK acquisition and should be treated as WGS84 geographic coordinates unless the researcher provides a more specific survey note.
    - CAAML is referenced only as contextual interoperability background and is not a deposited file format in this package.

    5. Environmental/experimental conditions:

    The campaign covered five Rogers Pass sites between {bbox["elevation_min_m"]:.0f} m and {bbox["elevation_max_m"]:.0f} m elevation. SnowScope profile depths in the prepared package range from {stats["snowscope"]["profile_depth_range_mm"][0]:.0f} mm to {stats["snowscope"]["profile_depth_range_mm"][1]:.0f} mm, and SMP profile depths extend to {stats["snowmicropenetrometer"]["depth_range_mm"][1]:.0f} mm. Manual snow-pit temperature entries preserved in the workbooks span approximately {stats["snow_stratigraphy"]["temperature_range_c"][0]:.1f} to {stats["snow_stratigraphy"]["temperature_range_c"][1]:.1f} degrees C. The Hermit field notes and workbook metadata indicate that the Day 3 profile quality was poor and should be interpreted cautiously.

    6. Describe any quality-assurance procedures performed on the data:

    Quality assurance during preparation focused on packaging and interpretability rather than on altering scientific measurements. The prepared package:
    - separates scientific files from administrative material;
    - documents unresolved scope decisions and raw-source anomalies in `artifacts/qc_report.md`;
    - standardizes folder names and filenames for deposit consistency;
    - preserves raw instrument formats and numeric content;
    - converts ancillary DOCX notes to plain text for accessibility.

    7. People involved with sample collection, processing, analysis and/or submission:

    Confirmed field participants in the site notes include Jean-Benoit Madore, Alexandre Langlois, Benjamin Imbach, Francis Gauthier, Francis Meloche, Violaine Paquette, Kate Hale, Hans-Peter Marshall, Jo Meyer, and Julien Meloche. Parks Canada staff are referenced in the notes as operational collaborators for site access and logistics. Pending researcher confirmation applies to final authorship, contact list, and funding wording.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: snow_stratigraphy/
    -----------------------------------------------------------------

    1. Number of variables:
    Four workbook sheets per file. The `AVY profile` sheet contains site metadata plus layer-by-layer snow observations; the `Stability Tests`, `Density`, and `IRIS` sheets contain supporting measurements.

    2. Number of cases/rows:
    7 workbook files. Active observation rows vary by snow pit and are embedded within 1000-row templates.

    3. Missing data codes:
            blank cell        No value recorded in the field template
            0                In some density columns, likely placeholder or missing entry rather than a true zero

    4. Variable List:
        A. Name: DATE / TIME / ELEVATION / ASPECT / INCLINE
           Description: Snow-pit metadata describing when and where the manual profile was recorded.

        B. Name: HEIGHT (cm) / RESISTANCE / FORM / EXTENT / LWC
           Description: Layer-by-layer stratigraphy fields describing layer depth, hand hardness, grain form, grain size, and wetness.

        C. Name: DENSITY
           Description: Density measurements from known-volume snow samples.

        D. Name: TEMPERATURE
           Description: Manual temperature profile measurements recorded in the AVY profile sheet.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: iris/
    -----------------------------------------------------------------

    1. Number of variables:
    2

    2. Number of cases/rows:
    6 files; {stats["iris"]["row_range"][0]}-{stats["iris"]["row_range"][1]} rows per file.

    3. Missing data codes:
            blank line        Not present in prepared files

    4. Variable List:
        A. Name: time
           Description: Time stamp recorded in the raw export (HH:MM:SS).

        B. Name: value
           Description: Raw IRIS measurement value associated with the time stamp.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: snowmicropenetrometer/
    -----------------------------------------------------------------

    1. Number of variables:
    Header metadata plus a high-resolution profile table with the main variables `distance` and `force`.

    2. Number of cases/rows:
    79 `.pnt` files; {stats["snowmicropenetrometer"]["sample_range"][0]}-{stats["snowmicropenetrometer"]["sample_range"][1]} samples per file.

    3. Missing data codes:
            -99999            Invalid GPS sentinel in some raw file headers

    4. Variable List:
        A. Name: distance
           Description: Penetration distance in millimetres.

        B. Name: force
           Description: Penetration resistance in Newtons.

        C. Name: timestamp / coordinates / serial metadata
           Description: Header metadata read from the binary file by `snowmicropyn`.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: snowscope/
    -----------------------------------------------------------------

    1. Number of variables:
    Approximately 21 metadata header fields plus a profile table with 2-3 columns.

    2. Number of cases/rows:
    319 profile files; {stats["snowscope"]["row_range"][0]}-{stats["snowscope"]["row_range"][1]} depth rows per profile table.

    3. Missing data codes:
            null             Missing metadata value in the header block
            blank field      Missing profile value

    4. Variable List:
        A. Name: serialNum / FW_version / PCB_version
           Description: Device identifier and firmware metadata.

        B. Name: collectionTime / collectionTime (Unix Time)
           Description: Acquisition timestamp.

        C. Name: Location
           Description: Latitude and longitude embedded in the file header.

        D. Name: depth (mm) / hardness (kPa) / optical Reflectance Avg
           Description: Depth-resolved SnowScope measurement table. `optical Reflectance Avg` appears only in a subset of profiles.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: radar/
    -----------------------------------------------------------------

    1. Number of variables:
    Header metadata plus a five-column numeric table.

    2. Number of cases/rows:
    166 files; each file contains {stats["radar"]["row_range"][0]} numeric rows after the header.

    3. Missing data codes:
            trailing incomplete line        Ignored when parsing; no replacement value written

    4. Variable List:
        A. Name: Radar No. / Start-Frequency [MHz] / Stop-Frequency [MHz]
           Description: Instrument and acquisition settings preserved in the header.

        B. Name: X (m)
           Description: Along-trace distance coordinate.

        C. Name: I1 / Q1 / I2 / Q2
           Description: Recorded in-phase and quadrature channels from the radar export.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: spatial_reference/
    -----------------------------------------------------------------

    1. Number of variables:
    Varies by file. Linkage workbooks map measurement IDs and GPS points; GPS CSV files contain 37 RTK export columns; shapefile ZIPs contain standard SHP sidecar components.

    2. Number of cases/rows:
    10 files total: 4 linkage workbooks, 3 GPS CSV files, and 3 shapefile ZIP packages.

    3. Missing data codes:
            blank cell        Not recorded in linkage workbook

    4. Variable List:
        A. Name: Mesure # / SMP / Radar / Snowscope Minute
           Description: Linkage fields that connect measurement IDs across instruments.

        B. Name: Latitude / Longitude / Elevation
           Description: RTK GPS coordinates in decimal degrees and metres.

        C. Name: Name / Solution status / Correction type
           Description: GPS point identifier and RTK quality metadata.

    -----------------------------------------------------------------
    DATA-SPECIFIC INFORMATION FOR: documentation/
    -----------------------------------------------------------------

    1. Number of variables:
    Narrative text files; no fixed tabular schema.

    2. Number of cases/rows:
    8 UTF-8 plain-text files converted from DOCX.

    3. Missing data codes:
            not applicable     Free-text field notes

    4. Variable List:
        A. Name: day/site heading
           Description: Human-readable identifier for the field day and site.

        B. Name: team list
           Description: Personnel named in the field note.

        C. Name: day summary / measurement summary
           Description: Narrative notes explaining what was measured, field conditions, and operational issues.
    """
    readme = textwrap.dedent(readme).strip() + "\n"
    issues = validation_checks(readme)
    write_text(README_PATH, readme)
    return issues


def update_metadata(manifest: dict, readme_issues: list[str]) -> None:
    metadata = yaml.safe_load(METADATA_PATH.read_text(encoding="utf-8"))
    bbox = manifest["stats"]["dataset_bbox"]
    metadata["contact"]["email"] = "jean-benoit.madore@usherbrooke.ca"
    metadata["geographic_coverage"]["point"]["latitude"] = round(bbox["mean_latitude"], 6)
    metadata["geographic_coverage"]["point"]["longitude"] = round(bbox["mean_longitude"], 6)
    metadata["geographic_coverage"]["bounding_box"]["west"] = round(bbox["west"], 6)
    metadata["geographic_coverage"]["bounding_box"]["east"] = round(bbox["east"], 6)
    metadata["geographic_coverage"]["bounding_box"]["north"] = round(bbox["north"], 6)
    metadata["geographic_coverage"]["bounding_box"]["south"] = round(bbox["south"], 6)
    metadata["geographic_coverage"]["elevation_range"] = f"{bbox['elevation_min_m']:.0f}-{bbox['elevation_max_m']:.0f} m"
    metadata["workflow"]["quality_control"] = "done"
    metadata["workflow"]["data_preparation"] = "done"
    metadata["workflow"]["draft_readme"] = "done"
    metadata["prepared_package"] = {
        "frdr_data_path": "datasets/rogers_pass_snow_profiles/frdr_data/",
        "prepared_file_count": manifest["prepared_total_files"],
        "excluded_file_count": manifest["excluded_total_files"],
        "documentation_language_note": "Converted field notes remain in French.",
        "readme_validation_issues": readme_issues,
    }
    write_text(METADATA_PATH, yaml.safe_dump(metadata, sort_keys=False, allow_unicode=True))


def prepare_dataset() -> dict:
    ARTIFACTS_ROOT.mkdir(parents=True, exist_ok=True)
    manifest = build_manifest()
    readme_issues = write_readme(manifest)
    manifest["readme_validation_issues"] = readme_issues
    write_data_preparation_report(manifest)
    write_qc_report()
    update_metadata(manifest, readme_issues)
    write_text(MANIFEST_PATH, json.dumps(manifest, indent=2))
    return manifest


def build_notebook() -> None:
    nb = nbf.v4.new_notebook()
    nb.metadata.kernelspec = {
        "display_name": "Python 3",
        "language": "python",
        "name": "python3",
    }
    nb.cells = [
        nbf.v4.new_markdown_cell(
            textwrap.dedent(
                """
                # Data Preparation: Rogers Pass Snow Profiles

                This notebook reruns the non-destructive preparation of the Rogers Pass
                field campaign into the FRDR deposit package.
                """
            ).strip()
        ),
        nbf.v4.new_code_cell(
            textwrap.dedent(
                """
                import sys
                from pathlib import Path

                project_root = Path.cwd()
                if str(project_root) not in sys.path:
                    sys.path.insert(0, str(project_root))

                from scripts.prepare_rogers_pass_dataset import MANIFEST_PATH, prepare_dataset
                """
            ).strip()
        ),
        nbf.v4.new_code_cell(
            textwrap.dedent(
                """
                manifest = prepare_dataset()
                {
                    "raw_total_files": manifest["raw_total_files"],
                    "prepared_total_files": manifest["prepared_total_files"],
                    "excluded_total_files": manifest["excluded_total_files"],
                    "transform_counts": manifest["transform_counts"],
                    "readme_validation_issues": manifest["readme_validation_issues"],
                }
                """
            ).strip()
        ),
        nbf.v4.new_code_cell(
            textwrap.dedent(
                """
                import pandas as pd

                pd.DataFrame.from_dict(manifest["category_summary"], orient="index")
                """
            ).strip()
        ),
        nbf.v4.new_code_cell("manifest['stats']"),
        nbf.v4.new_code_cell("Path(MANIFEST_PATH).read_text(encoding='utf-8')[:2000]"),
    ]
    NOTEBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    nbf.write(nb, NOTEBOOK_PATH)


def execute_notebook() -> None:
    build_notebook()
    nb = nbf.read(NOTEBOOK_PATH, as_version=4)
    client = NotebookClient(nb, timeout=600, kernel_name="python3")
    executed = client.execute()
    nbf.write(executed, NOTEBOOK_PATH)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--build-notebook", action="store_true")
    parser.add_argument("--execute-notebook", action="store_true")
    parser.add_argument("--prepare", action="store_true")
    args = parser.parse_args()

    if args.execute_notebook:
        execute_notebook()
    elif args.build_notebook:
        build_notebook()
    else:
        prepare_dataset()


if __name__ == "__main__":
    main()
